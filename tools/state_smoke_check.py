from __future__ import annotations

import shutil
import os
import sys
import uuid
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

import main as psl_main  # noqa: E402
import generate_license as license_generator  # noqa: E402


MIXED_LABEL = "\u0421\u043c\u0435\u0448\u0430\u043d\u043d\u0430\u044f"
CASH_LABEL = "\u041d\u0430\u043b\u0438\u0447\u043d\u044b\u0435"
CARD_LABEL = "\u041a\u0430\u0440\u0442\u0430"


def ensure(condition: bool, message: str) -> None:
    if not condition:
        raise AssertionError(message)


def decode_json(response) -> dict:
    return response.get_json(silent=True) or {}


def backup_file(path: Path) -> bytes | None:
    return path.read_bytes() if path.exists() else None


def restore_file(path: Path, original: bytes | None) -> None:
    if path.exists():
        path.unlink()
    if original is not None:
        path.write_bytes(original)


def build_state_payload() -> dict:
    return {
        "schemaVersion": psl_main.SCHEMA_VERSION,
        "version": 1,
        "lastModified": 1_900_000_000_000,
        "stations": [
            {
                "id": 5,
                "name": "Arcade Racer",
                "stationType": "simulator",
                "status": "idle",
                "startTime": None,
                "endTime": None,
                "tariffId": None,
                "history": [],
                "activeSessionId": None,
                "lastClosedSnapshot": None,
            },
            {
                "id": 2,
                "name": "PS2 Renamed",
                "stationType": "ps",
                "status": "idle",
                "startTime": None,
                "endTime": None,
                "tariffId": None,
                "history": [],
                "activeSessionId": None,
                "lastClosedSnapshot": None,
            },
        ],
        "sessions": {
            "2026-04-29": [
                {
                    "id": "sess-1",
                    "stationId": 5,
                    "stationName": "Old Legacy Sim",
                    "stationType": "ps",
                    "startTime": 1714377600000,
                    "endTime": 1714381200000,
                    "tariffId": "sim_1",
                    "tariffLabel": "1 час",
                    "mode": "manual",
                    "totalAmount": 600,
                    "paymentMethod": "cash",
                    "sales": [
                        {
                            "id": "sale-1",
                            "time": 1714377600000,
                            "type": "start_tariff",
                            "label": "Старт",
                            "minutes": 30,
                            "amount": 300,
                            "paymentMethod": "cash",
                        },
                        {
                            "id": "sale-2",
                            "time": 1714379400000,
                            "type": "extend_tariff",
                            "label": "Продление",
                            "minutes": 30,
                            "amount": 300,
                            "paymentMethod": "card",
                        },
                    ],
                }
            ]
        },
        "settings": {
            "graceMinutes": 10,
            "overdueMinutes": 30,
            "customRates": {"ps": 600, "simulator": 900, "switch": 500},
            "stationDefinitions": [
                {"id": 5, "name": "Arcade Racer", "type": "simulator"},
                {"id": 2, "name": "PS2 Renamed", "type": "ps"},
            ],
            "tariffGroups": {
                "ps": [{"id": "ps_1", "label": "1 час", "minutes": 60, "price": 600}],
                "simulator": [{"id": "sim_1", "label": "1 час", "minutes": 60, "price": 900}],
                "switch": [{"id": "sw_1", "label": "1 час", "minutes": 60, "price": 500}],
            },
            "notificationSound": False,
            "pinHash": "fnv1a_test",
        },
        "achievements": {
            "version": 1,
            "backfillVersion": 3,
            "unlocked": {},
            "progress": {"full_house": 4},
            "dynamicTargets": {"full_house": 6},
            "unseenIds": [],
            "breakEvents": [],
            "cleanTrackingStartedAt": 1714300000000,
        },
    }


def assert_report_workbook_content(workbook) -> None:
    summary = workbook.worksheets[0]
    sessions = workbook.worksheets[1]
    sales = workbook.worksheets[2]

    ensure(summary["A7"].value == 600, f"report summary revenue should be 600, got {summary['A7'].value}")
    ensure(summary["C7"].value == 1, f"report summary sessions count should be 1, got {summary['C7'].value}")
    ensure(summary["B12"].value == 600, f"report summary simulator total should be 600, got {summary['B12'].value}")
    ensure(summary["B17"].value == 300, f"report summary cash total should be 300, got {summary['B17'].value}")
    ensure(summary["B18"].value == 300, f"report summary card total should be 300, got {summary['B18'].value}")

    ensure(sessions["H2"].value == 600, f"report sessions total should be 600, got {sessions['H2'].value}")
    ensure(sessions["I2"].value == MIXED_LABEL, f"report sessions payment should be mixed, got {sessions['I2'].value}")

    ensure(sales["G2"].value == 300 and sales["H2"].value == CASH_LABEL, "report first sale row should be 300 cash")
    ensure(sales["G3"].value == 300 and sales["H3"].value == CARD_LABEL, "report second sale row should be 300 card")


def assert_filtered_report_workbook_content(workbook) -> None:
    summary = workbook.worksheets[0]
    sales = workbook.worksheets[2]
    data_rows = [
        [sales.cell(row, col).value for col in range(1, 9)]
        for row in range(2, sales.max_row + 1)
        if any(sales.cell(row, col).value is not None for col in range(1, 9))
    ]

    ensure(summary["A7"].value == 300, f"filtered report revenue should be 300, got {summary['A7'].value}")
    ensure(summary["C7"].value == 1, f"filtered report sessions count should stay 1, got {summary['C7'].value}")
    ensure(summary["B17"].value == 0, f"filtered report cash total should be 0, got {summary['B17'].value}")
    ensure(summary["B18"].value == 300, f"filtered report card total should be 300, got {summary['B18'].value}")
    ensure(len(data_rows) == 1, f"filtered report should have one non-empty sale row, got {len(data_rows)}")
    ensure(data_rows[0][7] == CARD_LABEL, f"filtered report sale payment should be card, got {data_rows[0][7]}")


def assert_today_workbook_content(workbook) -> None:
    shift = workbook.worksheets[0]
    sessions = workbook.worksheets[1]
    sales = workbook.worksheets[2]

    ensure(shift["B6"].value == 600, f"today revenue should be 600, got {shift['B6'].value}")
    ensure(shift["B7"].value == 1, f"today sessions count should be 1, got {shift['B7'].value}")
    ensure(shift["B18"].value == 300, f"today cash total should be 300, got {shift['B18'].value}")
    ensure(shift["B19"].value == 300, f"today card total should be 300, got {shift['B19'].value}")

    ensure(sessions["H2"].value == 600, f"today sessions total should be 600, got {sessions['H2'].value}")
    ensure(sessions["I2"].value == MIXED_LABEL, f"today sessions payment should be mixed, got {sessions['I2'].value}")

    ensure(sales["G2"].value == 300 and sales["H2"].value == CASH_LABEL, "today first sale row should be 300 cash")
    ensure(sales["G3"].value == 300 and sales["H3"].value == CARD_LABEL, "today second sale row should be 300 card")


def main() -> int:
    checks: list[str] = []
    client = psl_main.app.test_client()

    license_path = Path(psl_main._license_file_path())
    state_path = psl_main._app_dir() / psl_main.STATE_FILE
    state_bak_path = psl_main._app_dir() / psl_main.STATE_BAK_FILE
    history_dir = psl_main._app_dir() / psl_main.STATE_HISTORY_DIR

    original_license = backup_file(license_path)
    original_state = backup_file(state_path)
    original_state_bak = backup_file(state_bak_path)
    original_retention = os.environ.get("PS_LOUNGE_RETENTION_DAYS")
    history_backup_dir = ROOT / ".tmp_state_history_backup"

    try:
        os.environ["PS_LOUNGE_RETENTION_DAYS"] = "9999"
        if history_backup_dir.exists():
            shutil.rmtree(history_backup_dir, ignore_errors=True)
        if history_dir.exists():
            shutil.copytree(history_dir, history_backup_dir)
            shutil.rmtree(history_dir, ignore_errors=True)

        if license_path.exists():
            license_path.unlink()
        if state_path.exists():
            state_path.unlink()
        if state_bak_path.exists():
            state_bak_path.unlink()

        fingerprint = psl_main._machine_fingerprint()
        token = license_generator.build_token(
            "State Smoke Customer",
            fingerprint,
            str(uuid.uuid4()),
            seat_type="single_device",
            product="pslounge-desktop",
        )
        activated = psl_main._activate_license_token(token)
        ensure(activated.get("licensed") is True, "test license activation failed")
        checks.append("license")

        state_payload = build_state_payload()
        response = client.post("/api/backup", json=state_payload)
        payload = decode_json(response)
        ensure(response.status_code == 200 and payload.get("ok") is True, f"backup write failed: {response.status_code} {payload}")
        checks.append("backup-write")

        response = client.get("/api/backup")
        payload = decode_json(response)
        ensure(response.status_code == 200, f"backup read returned {response.status_code}")
        ensure(payload.get("lastModified") == state_payload["lastModified"], "backup read lost lastModified")
        ensure(payload.get("achievements", {}).get("dynamicTargets", {}).get("full_house") == 6, "dynamicTargets.full_house was not preserved")
        ensure(payload.get("settings", {}).get("stationDefinitions", [])[0].get("id") == 5, "station definition order was not preserved in backup payload")
        checks.append("backup-read")

        disk_state = psl_main._read_state_file()
        ensure(disk_state.get("source") == "primary", f"expected primary state source, got {disk_state.get('source')}")
        ensure(disk_state.get("recoveryDiagnostics", {}).get("selectedSource") == "primary", "recovery diagnostics should report primary source")
        checks.append("disk-read")

        session_rows, sales_rows, _ = psl_main._collect_report_rows(disk_state, "2026-04-29", "2026-04-29")
        ensure(
            session_rows and session_rows[0]["stationName"] == "Arcade Racer",
            f"report rows should use station name from saved definitions, got {session_rows[0]['stationName'] if session_rows else '<none>'}",
        )
        ensure(session_rows and session_rows[0]["stationType"] == "Симулятор гонок", "report rows should use station type from saved definitions")
        ensure(session_rows and session_rows[0]["_paymentRaw"] == "mixed", "session rows should preserve mixed payment state")
        ensure(sales_rows and sales_rows[0]["stationName"] == "Arcade Racer", "sales rows should use station name from saved definitions")
        checks.append("report-metadata")

        response = client.get("/api/export/report.xlsx?from=2026-04-29&to=2026-04-29")
        ensure(response.status_code == 200, f"report export returned {response.status_code}")
        content_type = response.headers.get("Content-Type", "")
        ensure("spreadsheetml" in content_type or "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in content_type, "report export content-type mismatch")
        assert_report_workbook_content(load_workbook(BytesIO(response.data), data_only=True))

        response = client.get("/api/export/report.xlsx?from=2026-04-29&to=2026-04-29&payment=card")
        ensure(response.status_code == 200, f"filtered report export returned {response.status_code}")
        assert_filtered_report_workbook_content(load_workbook(BytesIO(response.data), data_only=True))
        checks.append("report-export-content")

        response = client.get("/api/export/today.xlsx?date=2026-04-29")
        ensure(response.status_code == 200, f"today export returned {response.status_code}")
        assert_today_workbook_content(load_workbook(BytesIO(response.data), data_only=True))
        checks.append("today-export-content")

        history_snapshot = history_dir / f"{psl_main.datetime.now().strftime('%Y-%m-%d')}.json"
        ensure(history_snapshot.exists(), "daily state history snapshot was not created")
        checks.append("history-snapshot")

    finally:
        if original_retention is None:
            os.environ.pop("PS_LOUNGE_RETENTION_DAYS", None)
        else:
            os.environ["PS_LOUNGE_RETENTION_DAYS"] = original_retention
        restore_file(license_path, original_license)
        restore_file(state_path, original_state)
        restore_file(state_bak_path, original_state_bak)
        if history_dir.exists():
            shutil.rmtree(history_dir, ignore_errors=True)
        if history_backup_dir.exists():
            shutil.copytree(history_backup_dir, history_dir)
            shutil.rmtree(history_backup_dir, ignore_errors=True)

    print(f"[ok] state smoke check passed: {', '.join(checks)}")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except AssertionError as exc:
        print(f"[fail] {exc}", file=sys.stderr)
        raise SystemExit(1)
