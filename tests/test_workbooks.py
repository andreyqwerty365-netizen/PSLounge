from io import BytesIO

import pytest
from openpyxl import load_workbook

from pslounge.workbooks import _build_report_workbook, _build_shift_workbook


@pytest.mark.parametrize("count", [1, 2, 6, 9])
def test_dynamic_station_tables_do_not_shift_report_kpis(count):
    state = {
        "settings": {
            "stationDefinitions": [
                {"id": idx, "name": f"Station {idx}", "type": "ps"}
                for idx in range(1, count + 1)
            ]
        }
    }
    sessions = [
        {
            "stationName": "Station 1",
            "stationType": "PlayStation",
            "date": "01.10.2026",
            "start": "10:00",
            "end": "11:00",
            "paidTime": "1 час",
            "finish": "11:00",
            "close": "вручную",
            "duration": "1 час",
            "sum": 300,
            "payment": "Наличные",
            "mode": "manual",
            "_endTime": 1,
            "_durationMinutes": 60,
        }
    ]
    sales = [
        {
            "stationName": "Station 1",
            "stationType": "PlayStation",
            "date": "01.10.2026",
            "time": "10:00",
            "type": "Старт",
            "operation": "Старт",
            "label": "1 час",
            "minutes": 60,
            "amount": 300,
            "payment": "Наличные",
        }
    ]
    report = _build_report_workbook(
        "2026-10-01", "2026-10-01", sessions, sales, ["2026-10-01"], state
    )
    stream = BytesIO()
    report.save(stream)
    sheet = load_workbook(BytesIO(stream.getvalue()), data_only=True).active
    assert sheet["A7"].value == 300
    assert sheet["C7"].value == 1
    assert sheet["A9"].value == "Средний чек"
    assert sheet["A13"].value == "Выручка по типам станций"
    assert [sheet.cell(8 + idx, 8).value for idx in range(count)] == [
        f"Station {idx}" for idx in range(1, count + 1)
    ]
    assert sheet["I8"].value == 1
    shift = _build_shift_workbook("2026-10-01", sessions, sales, state).active
    assert shift["B6"].value == 300
    assert [shift.cell(13 + idx, 1).value for idx in range(count)] == [
        f"Station {idx}" for idx in range(1, count + 1)
    ]
    cash_rows = [row for row in shift.iter_rows() if row[0].value == "Наличные"]
    assert cash_rows[0][1].value == 300


def test_station_names_are_excel_text_not_formulas():
    name = "=1+1"
    state = {
        "settings": {"stationDefinitions": [{"id": 1, "name": name, "type": "ps"}]}
    }
    for workbook in (
        _build_shift_workbook("2026-10-01", [], [], state),
        _build_report_workbook(
            "2026-10-01", "2026-10-01", [], [], ["2026-10-01"], state
        ),
    ):
        stream = BytesIO()
        workbook.save(stream)
        restored = load_workbook(BytesIO(stream.getvalue()), data_only=False)
        cells = [
            cell
            for sheet in restored
            for row in sheet.iter_rows()
            for cell in row
            if cell.value == name
        ]
        assert cells
        assert all(cell.data_type == "s" for cell in cells)
