# -*- coding: utf-8 -*-
from __future__ import annotations


def _excel_styles():
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter

    header_fill = PatternFill("solid", fgColor="1F2A44")
    subheader_fill = PatternFill("solid", fgColor="EAF0FF")
    section_fill = PatternFill("solid", fgColor="F3F6FB")
    total_fill = PatternFill("solid", fgColor="EEF2F8")
    header_font = Font(bold=True, color="FFFFFF")
    title_font = Font(bold=True, size=16)
    subtitle_font = Font(bold=True, size=12)
    bold_font = Font(bold=True)
    thin = Side(style="thin", color="D9DEE8")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="left", vertical="center")
    left = Alignment(horizontal="left", vertical="center")

    def style_header(ws, row_idx):
        for cell in ws[row_idx]:
            if cell.value in (None, ""):
                continue
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center
            cell.border = border

    def style_range(
        ws, min_row, max_row, min_col, max_col, fill=None, font=None, alignment=None
    ):
        for row in ws.iter_rows(
            min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col
        ):
            for cell in row:
                if fill:
                    cell.fill = fill
                if font:
                    cell.font = font
                if alignment:
                    cell.alignment = alignment
                cell.border = border

    def auto_width(ws, min_width=10, max_width=32):
        dims = {}
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                val = str(cell.value)
                dims[cell.column] = max(dims.get(cell.column, 0), len(val))
                if cell.border is None:
                    cell.border = border
        for idx, length in dims.items():
            ws.column_dimensions[get_column_letter(idx)].width = min(
                max(length + 2, min_width), max_width
            )

    return {
        "header_fill": header_fill,
        "subheader_fill": subheader_fill,
        "section_fill": section_fill,
        "total_fill": total_fill,
        "header_font": header_font,
        "title_font": title_font,
        "subtitle_font": subtitle_font,
        "bold_font": bold_font,
        "border": border,
        "center": center,
        "left": left,
        "style_header": style_header,
        "style_range": style_range,
        "auto_width": auto_width,
    }


EXCEL_RUB_FORMAT = '# ##0" ₽"'


EXCEL_INT_FORMAT = "0"


def _apply_number_format_range(
    ws, col_idx: int, start_row: int, end_row: int, fmt: str
):
    if end_row < start_row:
        return
    for row in range(start_row, end_row + 1):
        cell = ws.cell(row=row, column=col_idx)
        if cell.value in (None, ""):
            continue
        if isinstance(cell.value, (int, float)):
            cell.number_format = fmt


def _add_excel_table(
    ws,
    table_name: str,
    start_row: int,
    data_end_row: int,
    end_col: int,
    totals: dict[int, str] | None = None,
):
    from openpyxl.worksheet.table import Table, TableStyleInfo, TableColumn
    from openpyxl.utils import get_column_letter

    table_end_row = data_end_row + (1 if totals else 0)
    ref = f"A{start_row}:{get_column_letter(end_col)}{table_end_row}"
    tab = Table(displayName=table_name, ref=ref)
    tab.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )

    headers = [
        ws.cell(start_row, col_idx).value or f"Column{col_idx}"
        for col_idx in range(1, end_col + 1)
    ]
    tab.tableColumns = [
        TableColumn(id=idx, name=str(name)) for idx, name in enumerate(headers, start=1)
    ]
    if totals:
        tab.totalsRowCount = 1
        tab.totalsRowShown = True
        for col_idx, func in totals.items():
            if 1 <= col_idx <= len(tab.tableColumns):
                tab.tableColumns[col_idx - 1].totalsRowFunction = func
    ws.add_table(tab)
    return table_end_row


def _write_data_sheet(
    ws,
    headers: list[str],
    rows: list[list],
    table_name: str,
    money_cols: tuple[int, ...] = (),
    int_cols: tuple[int, ...] = (),
    totals: dict[int, str] | None = None,
):
    S = _excel_styles()

    ws.append(headers)
    S["style_header"](ws, 1)

    for row in rows:
        ws.append(row)

    data_end_row = ws.max_row
    table_end_row = data_end_row
    if rows:
        table_end_row = _add_excel_table(
            ws, table_name, 1, data_end_row, len(headers), totals=totals
        )

    start_fmt_row = 2
    end_fmt_row = data_end_row

    for col_idx in money_cols:
        _apply_number_format_range(
            ws, col_idx, start_fmt_row, end_fmt_row, EXCEL_RUB_FORMAT
        )
        if totals and rows:
            ws.cell(row=table_end_row, column=col_idx).number_format = EXCEL_RUB_FORMAT

    for col_idx in int_cols:
        _apply_number_format_range(
            ws, col_idx, start_fmt_row, end_fmt_row, EXCEL_INT_FORMAT
        )
        if totals and rows:
            ws.cell(row=table_end_row, column=col_idx).number_format = EXCEL_INT_FORMAT

    S["auto_width"](ws)
    from openpyxl.styles import Alignment

    for row in ws.iter_rows(
        min_row=2, max_row=ws.max_row, min_col=1, max_col=len(headers)
    ):
        for cell in row:
            cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.freeze_panes = "A2"


def _add_revenue_chart(ws, min_row: int, max_row: int, anchor: str = "D12"):
    from openpyxl.chart import BarChart, Reference
    from openpyxl.chart.label import DataLabelList

    if max_row <= min_row:
        return

    chart = BarChart()
    chart.type = "col"
    chart.style = 11
    chart.title = "Выручка по дням"
    chart.height = 7.2
    chart.width = 12.8
    chart.legend = None
    chart.gapWidth = 65
    chart.overlap = 0
    chart.varyColors = False

    chart.y_axis.title = "Выручка, ₽"
    chart.y_axis.number_format = "# ##0"
    try:
        chart.y_axis.majorGridlines = None
    except Exception:
        pass
    try:
        chart.x_axis.majorGridlines = None
    except Exception:
        pass
    chart.x_axis.tickLblPos = "low"
    chart.x_axis.delete = False
    try:
        chart.x_axis.title = None
    except Exception:
        pass

    data = Reference(ws, min_col=2, min_row=min_row, max_row=max_row)
    cats = Reference(ws, min_col=1, min_row=min_row + 1, max_row=max_row)

    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)

    if chart.series:
        series = chart.series[0]
        series.graphicalProperties.line.noFill = True
        try:
            series.graphicalProperties.solidFill = "4F81BD"
        except Exception:
            pass

    chart.dLbls = DataLabelList()
    chart.dLbls.showVal = True
    chart.dLbls.showCatName = False
    chart.dLbls.showLegendKey = False
    chart.dLbls.showSerName = False
    chart.dLbls.showPercent = False
    chart.dLbls.position = "outEnd"

    ws.add_chart(chart, anchor)


def _set_summary_value(cell, value, kind: str):
    cell.value = value
    if kind == "money":
        cell.number_format = EXCEL_RUB_FORMAT
    elif kind == "int":
        cell.number_format = EXCEL_INT_FORMAT


def _apply_daily_heatmap(ws, col_idx: int, start_row: int, end_row: int):
    from openpyxl.styles import PatternFill, Font

    values = []
    for row in range(start_row, end_row + 1):
        val = ws.cell(row=row, column=col_idx).value
        if isinstance(val, (int, float)):
            values.append(float(val))
    if not values:
        return
    vmax = max(values) or 1
    for row in range(start_row, end_row + 1):
        cell = ws.cell(row=row, column=col_idx)
        val = cell.value if isinstance(cell.value, (int, float)) else 0
        ratio = max(0.0, min(1.0, float(val) / vmax if vmax else 0.0))
        if ratio >= 0.85:
            fill = PatternFill("solid", fgColor="C6EFCE")
        elif ratio >= 0.55:
            fill = PatternFill("solid", fgColor="DDEBF7")
        elif ratio > 0:
            fill = PatternFill("solid", fgColor="FFF2CC")
        else:
            fill = PatternFill("solid", fgColor="F3F6FB")
        cell.fill = fill
        if ratio >= 0.85:
            cell.font = Font(bold=True)


def _add_station_load_chart(ws, min_row: int, max_row: int, anchor: str = "G12"):
    from openpyxl.chart import BarChart, Reference
    from openpyxl.chart.label import DataLabelList

    if max_row <= min_row:
        return

    chart = BarChart()
    chart.type = "bar"
    chart.style = 10
    chart.title = "Загрузка станций"
    chart.height = 6.4
    chart.width = 8.6
    chart.legend = None
    chart.gapWidth = 40
    chart.varyColors = False

    try:
        chart.y_axis.majorGridlines = None
    except Exception:
        pass
    try:
        chart.x_axis.majorGridlines = None
    except Exception:
        pass
    chart.x_axis.title = "Часы"
    chart.x_axis.number_format = "0.0"
    chart.y_axis.title = None

    data = Reference(ws, min_col=8, min_row=min_row, max_row=max_row)
    cats = Reference(ws, min_col=7, min_row=min_row + 1, max_row=max_row)

    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    if chart.series:
        series = chart.series[0]
        series.graphicalProperties.line.noFill = True
        try:
            series.graphicalProperties.solidFill = "6FA8DC"
        except Exception:
            pass

    chart.dLbls = DataLabelList()
    chart.dLbls.showVal = True
    chart.dLbls.position = "outEnd"
    ws.add_chart(chart, anchor)


def _literalize_report_text(wb):
    """Report strings are data, including names beginning with '='.

    Aggregation is represented by table totalsRowFunction metadata, not formulas
    supplied by users. Keep report text from becoming executable Excel formulas.
    """
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    cell.data_type = "s"
    return wb
