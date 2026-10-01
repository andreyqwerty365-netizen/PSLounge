from __future__ import annotations
from ..excel import (
    _literalize_report_text,
    _excel_styles,
    _set_summary_value,
    _write_data_sheet,
)
from ..formatting import _format_date_ru, _minutes_label
from ..report_data import _report_station_names


def _build_shift_workbook(
    day_key: str,
    session_rows: list[dict],
    sales_rows: list[dict],
    state: dict | None = None,
):
    from openpyxl import Workbook

    S = _excel_styles()
    wb = Workbook()
    ws_shift = wb.active
    ws_shift.title = "Смена"
    ws_sessions = wb.create_sheet("Сессии за день")
    ws_sales = wb.create_sheet("Начисления за день")

    # Sheet 1: daily operational shift summary
    ws_shift.merge_cells("A1:D1")
    ws_shift["A1"] = "PS Lounge — Смена"
    ws_shift["A1"].font = S["title_font"]
    ws_shift["A3"] = "Дата"
    ws_shift["B3"] = _format_date_ru(day_key)
    ws_shift["A3"].font = S["bold_font"]
    ws_shift["A5"] = "Показатель"
    ws_shift["B5"] = "Значение"
    S["style_header"](ws_shift, 5)

    closed_sessions = [r for r in session_rows if r.get("_endTime")]
    revenue = sum(r["amount"] for r in sales_rows)
    session_count = len(session_rows)
    avg_check = round(revenue / session_count) if session_count else 0
    avg_minutes = (
        round(
            sum(max(0, int(r.get("_durationMinutes") or 0)) for r in closed_sessions)
            / len(closed_sessions)
        )
        if closed_sessions
        else 0
    )
    summary_rows = [
        ("Выручка за день", revenue, "money"),
        ("Сессий", session_count, "int"),
        ("Средний чек", avg_check, "money"),
        ("Средняя длительность", _minutes_label(avg_minutes), "text"),
    ]
    row = 6
    for label, val, kind in summary_rows:
        ws_shift[f"A{row}"] = label
        ws_shift[f"A{row}"].font = S["bold_font"]
        _set_summary_value(ws_shift[f"B{row}"], val, kind)
        row += 1

    row += 1
    ws_shift[f"A{row}"] = "Выручка по станциям"
    ws_shift[f"A{row}"].font = S["subtitle_font"]
    row += 1
    ws_shift[f"A{row}"] = "Станция"
    ws_shift[f"B{row}"] = "Выручка"
    S["style_header"](ws_shift, row)
    row += 1

    station_order = _report_station_names(state, session_rows, sales_rows)
    station_totals = {name: 0 for name in station_order}
    for sale in sales_rows:
        station_totals[sale["stationName"]] = (
            station_totals.get(sale["stationName"], 0) + sale["amount"]
        )
    for name in station_order:
        ws_shift[f"A{row}"] = name
        _set_summary_value(ws_shift[f"B{row}"], station_totals.get(name, 0), "money")
        row += 1

    row += 1
    ws_shift[f"A{row}"] = "Выручка по оплате"
    ws_shift[f"A{row}"].font = S["subtitle_font"]
    row += 1
    ws_shift[f"A{row}"] = "Оплата"
    ws_shift[f"B{row}"] = "Сумма"
    S["style_header"](ws_shift, row)
    row += 1
    payment_totals = {"Наличные": 0, "Карта": 0, "Перевод": 0}
    for sale in sales_rows:
        payment_totals[sale["payment"]] = (
            payment_totals.get(sale["payment"], 0) + sale["amount"]
        )
    for label in ("Наличные", "Карта", "Перевод"):
        ws_shift[f"A{row}"] = label
        _set_summary_value(ws_shift[f"B{row}"], payment_totals.get(label, 0), "money")
        row += 1

    row += 1
    ws_shift[f"A{row}"] = "Активные сессии сейчас"
    ws_shift[f"A{row}"].font = S["subtitle_font"]
    row += 1
    ws_shift[f"A{row}"] = "Станция"
    ws_shift[f"B{row}"] = "Оплачено времени"
    ws_shift[f"C{row}"] = "Окончание"
    ws_shift[f"D{row}"] = "Оплачено"
    S["style_header"](ws_shift, row)
    row += 1
    active = [r for r in session_rows if not r.get("_endTime")]
    if active:
        for r in active:
            ws_shift[f"A{row}"] = r["stationName"]
            ws_shift[f"B{row}"] = r["paidTime"]
            ws_shift[f"C{row}"] = "—"
            _set_summary_value(ws_shift[f"D{row}"], r["sum"], "money")
            row += 1
    else:
        ws_shift[f"A{row}"] = "Нет активных сессий"
        row += 1

    S["style_range"](ws_shift, 3, ws_shift.max_row, 1, 4, alignment=S["left"])
    S["style_header"](ws_shift, 5)
    S["auto_width"](ws_shift)
    ws_shift.freeze_panes = "A6"

    # Sheet 2: sessions for day
    sess_headers = [
        "Дата",
        "Станция",
        "Тип станции",
        "Оплачено времени",
        "Старт",
        "Финиш",
        "Длительность",
        "Сумма",
        "Оплата",
        "Завершение",
    ]
    session_table_rows = [
        [
            r["date"],
            r["stationName"],
            r["stationType"],
            r["paidTime"],
            r["start"],
            r["finish"],
            r["duration"],
            r["sum"],
            r["payment"],
            r["close"],
        ]
        for r in session_rows
    ]
    _write_data_sheet(
        ws_sessions,
        headers=sess_headers,
        rows=session_table_rows,
        table_name="ShiftSessionsTable",
        money_cols=(8,),
        totals={8: "sum"},
    )

    # Sheet 3: sales for day
    sale_headers = [
        "Дата",
        "Время",
        "Станция",
        "Тип станции",
        "Операция",
        "Минуты",
        "Сумма",
        "Оплата",
    ]
    sales_table_rows = [
        [
            r["date"],
            r["time"],
            r["stationName"],
            r["stationType"],
            r["operation"],
            r["minutes"],
            r["amount"],
            r["payment"],
        ]
        for r in sales_rows
    ]
    _write_data_sheet(
        ws_sales,
        headers=sale_headers,
        rows=sales_table_rows,
        table_name="ShiftSalesTable",
        money_cols=(7,),
        int_cols=(6,),
        totals={7: "sum"},
    )

    return _literalize_report_text(wb)
