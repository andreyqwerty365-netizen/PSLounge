from __future__ import annotations
from datetime import datetime
from ..excel import (
    _literalize_report_text,
    EXCEL_INT_FORMAT,
    EXCEL_RUB_FORMAT,
    _add_revenue_chart,
    _apply_daily_heatmap,
    _excel_styles,
    _set_summary_value,
    _write_data_sheet,
)
from ..formatting import _format_date_ru, _minutes_label
from ..report_data import _report_station_names

from .audit import _write_audit_sheet
from .audit import _top_analytics


def _build_report_workbook(
    from_key: str,
    to_key: str,
    session_rows: list[dict],
    sales_rows: list[dict],
    period_keys: list[str],
    state: dict | None = None,
):
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill, Font, Alignment
    from openpyxl.formatting.rule import DataBarRule

    S = _excel_styles()
    wb = Workbook()
    ws_summary = wb.active
    ws_summary.title = "Сводка"
    ws_sessions = wb.create_sheet("Сессии")
    ws_sales = wb.create_sheet("Начисления")

    # Summary layout
    ws_summary.merge_cells("A1:B1")
    ws_summary["A1"] = "PS Lounge"
    ws_summary["A1"].font = Font(bold=True, size=18)
    ws_summary["A1"].alignment = Alignment(horizontal="left", vertical="center")

    ws_summary.merge_cells("A2:B2")
    ws_summary["A2"] = "Отчёт за период"
    ws_summary["A2"].font = Font(bold=True, size=12, color="666666")
    ws_summary["A2"].alignment = Alignment(horizontal="left", vertical="center")

    ws_summary["A4"] = "Период"
    ws_summary["B4"] = (
        _format_date_ru(from_key)
        if from_key == to_key
        else f"{_format_date_ru(from_key)} — {_format_date_ru(to_key)}"
    )
    ws_summary["A4"].font = S["bold_font"]
    # Проверка данных переносим вниз, чтобы не перегружать верхнюю часть

    closed_sessions = [r for r in session_rows if r.get("_endTime")]
    revenue = sum(r["amount"] for r in sales_rows)
    session_count = len(session_rows)
    avg_check = round(revenue / session_count) if session_count else 0
    avg_minutes = (
        round(
            sum(max(0, int(r.get("_durationMinutes") or 0)) for r in closed_sessions)
            / len(closed_sessions)
        )
        if closed_sessions
        else 0
    )

    # KPI cards
    kpi_fill = PatternFill("solid", fgColor="F7FAFF")
    kpi_value_font = Font(bold=True, size=14)
    kpi_title_font = Font(bold=True, size=11, color="44546A")
    kpis = [
        ("Выручка", revenue, "money"),
        ("Сессий", session_count, "int"),
        ("Средний чек", avg_check, "money"),
        ("Средняя длительность", _minutes_label(avg_minutes), "text"),
    ]
    kpi_rows = [(6, 1, 2), (6, 3, 4), (9, 1, 2), (9, 3, 4)]
    for (label, val, kind), (r, c1, c2) in zip(kpis, kpi_rows):
        ws_summary.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
        ws_summary.merge_cells(
            start_row=r + 1, start_column=c1, end_row=r + 1, end_column=c2
        )
        title_cell = ws_summary.cell(r, c1)
        value_cell = ws_summary.cell(r + 1, c1)
        title_cell.value = label
        title_cell.font = kpi_title_font
        value_cell.value = val
        value_cell.font = kpi_value_font
        if kind == "money":
            value_cell.number_format = EXCEL_RUB_FORMAT
        elif kind == "int":
            value_cell.number_format = EXCEL_INT_FORMAT
        title_cell.alignment = value_cell.alignment = Alignment(
            horizontal="left", vertical="center"
        )
        for rr in (r, r + 1):
            for cc in range(c1, c2 + 1):
                cell = ws_summary.cell(rr, cc)
                cell.fill = kpi_fill
                cell.border = S["border"]

    # Revenue by types
    ws_summary.merge_cells("A13:B13")
    ws_summary["A13"] = "Выручка по типам станций"
    ws_summary["A13"].font = S["subtitle_font"]
    ws_summary["A14"] = "Тип станции"
    ws_summary["B14"] = "Выручка"
    S["style_header"](ws_summary, 14)
    station_totals = {"PlayStation": 0, "Симулятор гонок": 0, "Nintendo Switch": 0}
    for sale in sales_rows:
        station_totals[sale["stationType"]] = (
            station_totals.get(sale["stationType"], 0) + sale["amount"]
        )
    row = 15
    for label in ("PlayStation", "Симулятор гонок", "Nintendo Switch"):
        ws_summary[f"A{row}"] = label
        _set_summary_value(ws_summary[f"B{row}"], station_totals.get(label, 0), "money")
        row += 1

    # Payment methods
    ws_summary.merge_cells("A19:B19")
    ws_summary["A19"] = "Выручка по оплате"
    ws_summary["A19"].font = S["subtitle_font"]
    ws_summary["A20"] = "Оплата"
    ws_summary["B20"] = "Сумма"
    S["style_header"](ws_summary, 20)
    row = 21
    payment_totals = {"Наличные": 0, "Карта": 0, "Перевод": 0}
    for sale in sales_rows:
        payment_totals[sale["payment"]] = (
            payment_totals.get(sale["payment"], 0) + sale["amount"]
        )
    for label in ("Наличные", "Карта", "Перевод"):
        ws_summary[f"A{row}"] = label
        _set_summary_value(ws_summary[f"B{row}"], payment_totals.get(label, 0), "money")
        row += 1

    # Daily revenue table + heatmap
    ws_summary.merge_cells("A25:B25")
    ws_summary["A25"] = "Тепловая карта выручки по дням"
    ws_summary["A25"].font = S["subtitle_font"]
    daily_header_row = 26
    ws_summary["A26"] = "Дата"
    ws_summary["B26"] = "Выручка"
    S["style_header"](ws_summary, 26)
    row = 27
    day_totals = {k: 0 for k in period_keys}
    for sale in sales_rows:
        try:
            iso_key = datetime.strptime(sale["date"], "%d.%m.%Y").strftime("%Y-%m-%d")
        except Exception:
            continue
        day_totals[iso_key] = day_totals.get(iso_key, 0) + sale["amount"]
    for key in period_keys:
        ws_summary[f"A{row}"] = _format_date_ru(key)
        _set_summary_value(ws_summary[f"B{row}"], day_totals.get(key, 0), "money")
        row += 1
    daily_end_row = row - 1
    _apply_daily_heatmap(ws_summary, 2, 27, daily_end_row)

    # Top analytics
    ws_summary["E6"] = "TOP-аналитика"
    ws_summary["E6"].font = S["subtitle_font"]
    ws_summary["E7"] = "Показатель"
    ws_summary["F7"] = "Значение"
    for ref in ("E7", "F7"):
        ws_summary[ref].fill = S["header_fill"]
        ws_summary[ref].font = S["header_font"]
        ws_summary[ref].alignment = S["center"]
        ws_summary[ref].border = S["border"]

    top_items = _top_analytics(session_rows, sales_rows)
    top_rows = [
        (top_items[0][0], top_items[0][1], top_items[0][2]),
        (top_items[1][0], top_items[1][1], top_items[1][2]),
        ("", "", "blank"),
        (top_items[2][0], top_items[2][1], top_items[2][2]),
        (top_items[3][0], top_items[3][1], top_items[3][2]),
        ("", "", "blank"),
        (top_items[4][0], top_items[4][1], top_items[4][2]),
        (top_items[5][0], top_items[5][1], top_items[5][2]),
    ]
    top_row = 8
    for label, val, kind in top_rows:
        if kind == "blank":
            top_row += 1
            continue
        ws_summary.cell(top_row, 5).value = label
        ws_summary.cell(top_row, 5).font = S["bold_font"]
        ws_summary.cell(top_row, 5).alignment = S["left"]
        _set_summary_value(ws_summary.cell(top_row, 6), val, kind)
        ws_summary.cell(top_row, 6).alignment = S["left"]
        top_row += 1

    # Load by stations (compact visual table)
    ws_summary["H6"] = "Загрузка станций"
    ws_summary["H6"].font = S["subtitle_font"]
    ws_summary["H7"] = "Станция"
    ws_summary["I7"] = "Часы"
    for ref in ("H7", "I7"):
        ws_summary[ref].fill = S["header_fill"]
        ws_summary[ref].font = S["header_font"]
        ws_summary[ref].alignment = S["center"]
        ws_summary[ref].border = S["border"]
    load_minutes = {}
    for r in session_rows:
        key = r.get("stationName") or "—"
        mins = max(0, int(r.get("_durationMinutes") or 0))
        load_minutes[key] = load_minutes.get(key, 0) + mins
    station_order = _report_station_names(state, session_rows, sales_rows)
    load_row = 8
    for name in station_order:
        ws_summary.cell(load_row, 8).value = name
        ws_summary.cell(load_row, 8).alignment = S["left"]
        hours = round(load_minutes.get(name, 0) / 60, 1)
        ws_summary.cell(load_row, 9).value = hours
        ws_summary.cell(load_row, 9).number_format = '0.0" ч"'
        ws_summary.cell(load_row, 9).alignment = S["left"]
        load_row += 1
    load_end_row = load_row - 1
    ws_summary.conditional_formatting.add(
        f"I8:I{load_end_row}",
        DataBarRule(
            start_type="num",
            start_value=0,
            end_type="max",
            end_value=0,
            color="5B9BD5",
            showValue=True,
        ),
    )

    # Visual cleanup
    S["style_range"](ws_summary, 13, 23, 1, 2, alignment=S["left"])
    S["style_range"](ws_summary, 7, top_row - 1, 5, 6, alignment=S["left"])
    S["style_range"](ws_summary, 7, load_end_row, 8, 9, alignment=S["left"])
    S["style_range"](ws_summary, 26, daily_end_row, 1, 2, alignment=S["left"])

    for col, width in {
        "A": 18,
        "B": 20,
        "C": 16,
        "D": 16,
        "E": 28,
        "F": 34,
        "G": 12,
        "H": 28,
        "I": 16,
        "J": 20,
        "K": 16,
    }.items():
        ws_summary.column_dimensions[col].width = width

    from openpyxl.styles import Alignment

    for row in ws_summary.iter_rows(
        min_row=1, max_row=ws_summary.max_row, min_col=1, max_col=11
    ):
        for cell in row:
            cell.alignment = Alignment(horizontal="left", vertical="center")
    for ref in ("E7", "F7", "H7", "I7", "A14", "B14", "A20", "B20", "A26", "B26"):
        ws_summary[ref].alignment = Alignment(horizontal="left", vertical="center")
    ws_summary.freeze_panes = "A6"
    _add_revenue_chart(ws_summary, daily_header_row, daily_end_row, anchor="E18")

    # Sessions sheet
    session_table_rows = [
        [
            r["date"],
            r["stationName"],
            r["stationType"],
            r["paidTime"],
            r["start"],
            r["finish"],
            r["duration"],
            r["sum"],
            r["payment"],
            r["close"],
        ]
        for r in session_rows
    ]
    _write_data_sheet(
        ws_sessions,
        headers=[
            "Дата",
            "Станция",
            "Тип станции",
            "Оплачено времени",
            "Старт",
            "Финиш",
            "Длительность",
            "Сумма",
            "Оплата",
            "Завершение",
        ],
        rows=session_table_rows,
        table_name="ReportSessionsTable",
        money_cols=(8,),
        totals={8: "sum"},
    )

    # Sales sheet
    sales_table_rows = [
        [
            r["date"],
            r["time"],
            r["stationName"],
            r["stationType"],
            r["operation"],
            r["minutes"],
            r["amount"],
            r["payment"],
        ]
        for r in sales_rows
    ]
    _write_data_sheet(
        ws_sales,
        headers=[
            "Дата",
            "Время",
            "Станция",
            "Тип станции",
            "Операция",
            "Минуты",
            "Сумма",
            "Оплата",
        ],
        rows=sales_table_rows,
        table_name="ReportSalesTable",
        money_cols=(7,),
        int_cols=(6,),
        totals={7: "sum"},
    )

    audit_errors, audit_warnings = _write_audit_sheet(wb, session_rows, sales_rows)

    ws_summary["A33"] = "Проверка данных"
    ws_summary["A33"].font = S["bold_font"]
    ws_summary["A34"] = (
        "Ошибок не обнаружено"
        if not (audit_errors or audit_warnings)
        else f"Ошибки: {audit_errors} | Предупреждения: {audit_warnings}"
    )
    ws_summary["A34"].font = S["bold_font"]

    return _literalize_report_text(wb)
