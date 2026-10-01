from __future__ import annotations
from ..excel import (
    _write_data_sheet,
)
from ..formatting import _minutes_label, _money_int, _safe_int


def _build_audit_rows(
    session_rows: list[dict], sales_rows: list[dict]
) -> tuple[list[list], int, int]:
    sales_by_session: dict[str, dict[str, int]] = {}
    orphan_sales = 0

    for sale in sales_rows:
        session_key = str(sale.get("_sessionKey") or "").strip()
        if not session_key:
            orphan_sales += 1
            continue
        bucket = sales_by_session.setdefault(
            session_key, {"amount": 0, "minutes": 0, "count": 0}
        )
        bucket["amount"] += _money_int(sale.get("amount"))
        bucket["minutes"] += max(0, _safe_int(sale.get("minutes"), 0))
        bucket["count"] += 1

    rows: list[list] = []
    error_count = 0
    warning_count = 0

    for row in session_rows:
        session_key = str(row.get("_sessionKey") or "")
        sale_info = sales_by_session.get(
            session_key, {"amount": 0, "minutes": 0, "count": 0}
        )
        session_sum = _money_int(row.get("sum"))
        sales_sum = int(sale_info["amount"])
        paid_minutes = max(0, _safe_int(row.get("_paidMinutes"), 0))
        sales_minutes = int(sale_info["minutes"])
        duration_minutes = max(0, _safe_int(row.get("_durationMinutes"), 0))
        close_mode = str(row.get("close") or "").strip().lower()
        flags = list(dict.fromkeys(row.get("_integrityFlags") or []))

        findings: list[tuple[str, str]] = []

        if sales_sum > 0 and session_sum != sales_sum:
            findings.append(("Ошибка", "Сумма сессии не совпадает с суммой начислений"))
        if sales_minutes > 0 and paid_minutes > 0 and paid_minutes != sales_minutes:
            findings.append(
                ("Ошибка", "Оплаченные минуты не совпадают с минутами начислений")
            )
        if session_sum < 0 or sales_sum < 0:
            findings.append(("Ошибка", "Обнаружена отрицательная сумма"))
        if paid_minutes < 0 or sales_minutes < 0:
            findings.append(("Ошибка", "Обнаружено отрицательное количество минут"))

        if row.get("_endTime") and session_sum > 0 and duration_minutes == 0:
            findings.append(
                ("Предупреждение", "Нулевая длительность при ненулевой сумме")
            )
        if (
            row.get("_endTime")
            and row.get("_startTime")
            and close_mode == "вручную"
            and 0 <= duration_minutes <= 1
            and session_sum > 0
        ):
            findings.append(
                ("Предупреждение", "Ручное закрытие почти сразу после старта")
            )
        if (
            row.get("_endTime")
            and paid_minutes > 0
            and duration_minutes > paid_minutes + 60
        ):
            findings.append(
                (
                    "Предупреждение",
                    "Фактическая длительность заметно больше оплаченного времени",
                )
            )

        if "sales_total_mismatch" in flags and not any(
            desc == "Сумма сессии не совпадает с суммой начислений"
            for _, desc in findings
        ):
            findings.append(("Ошибка", "Итог сессии не совпадает с raw totalAmount"))
        if "duration_adjusted" in flags and row.get("_endTime"):
            findings.append(
                ("Предупреждение", "Длительность была нормализована при экспорте")
            )

        seen: set[tuple[str, str]] = set()
        for level, description in findings:
            if (level, description) in seen:
                continue
            seen.add((level, description))
            if level == "Ошибка":
                error_count += 1
            else:
                warning_count += 1
            rows.append(
                [
                    level,
                    row.get("date"),
                    row.get("stationName"),
                    row.get("stationType"),
                    row.get("payment"),
                    session_sum,
                    sales_sum,
                    paid_minutes,
                    sales_minutes,
                    row.get("duration"),
                    description,
                ]
            )

    if orphan_sales:
        error_count += 1
        rows.append(
            [
                "Ошибка",
                "",
                "—",
                "—",
                "—",
                0,
                0,
                0,
                0,
                "—",
                f"Найдены начисления без привязки к сессии: {orphan_sales}",
            ]
        )

    return rows, error_count, warning_count


def _write_audit_sheet(
    wb, session_rows: list[dict], sales_rows: list[dict]
) -> tuple[int, int]:
    from openpyxl.styles import Alignment, PatternFill, Font

    rows, error_count, warning_count = _build_audit_rows(session_rows, sales_rows)
    if not rows:
        return 0, 0

    ws = wb.create_sheet("Audit")
    headers = [
        "Уровень",
        "Дата",
        "Станция",
        "Тип станции",
        "Оплата",
        "Сумма сессии",
        "Сумма начислений",
        "Оплачено минут",
        "Минут начислено",
        "Длительность",
        "Описание",
    ]
    _write_data_sheet(
        ws,
        headers=headers,
        rows=rows,
        table_name="AuditTable",
        money_cols=(6, 7),
        int_cols=(8, 9),
    )
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["K"].width = 55

    warn_fill = PatternFill("solid", fgColor="FFD966")
    err_fill = PatternFill("solid", fgColor="F4B183")
    bold_font = Font(bold=True)

    for row in range(2, ws.max_row + 1):
        level_cell = ws.cell(row=row, column=1)
        text = str(level_cell.value or "").strip()
        if text == "Предупреждение":
            level_cell.fill = warn_fill
        elif text == "Ошибка":
            level_cell.fill = err_fill
        level_cell.font = bold_font
        for col in range(1, 12):
            ws.cell(row=row, column=col).alignment = Alignment(
                horizontal="left", vertical="center"
            )
        ws.cell(row=row, column=11).alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
    return error_count, warning_count


def _top_analytics(
    session_rows: list[dict], sales_rows: list[dict]
) -> list[tuple[str, object, str]]:
    top_session = max(
        session_rows,
        key=lambda r: (
            _money_int(r.get("sum")),
            _safe_int(r.get("_paidMinutes"), 0),
            str(r.get("stationName") or ""),
        ),
        default=None,
    )
    closed_sessions = [r for r in session_rows if r.get("_endTime")]
    longest_session = max(
        closed_sessions,
        key=lambda r: (
            _safe_int(r.get("_durationMinutes"), 0),
            _money_int(r.get("sum")),
            str(r.get("stationName") or ""),
        ),
        default=None,
    )

    station_revenue: dict[str, int] = {}
    for sale in sales_rows:
        name = str(sale.get("stationName") or "—")
        station_revenue[name] = station_revenue.get(name, 0) + _money_int(
            sale.get("amount")
        )
    top_station_name = "—"
    top_station_revenue = 0
    if station_revenue:
        top_station_name, top_station_revenue = max(
            station_revenue.items(), key=lambda item: (item[1], item[0])
        )

    analytics: list[tuple[str, object, str]] = []
    if top_session:
        analytics.append(
            (
                "Самая дорогая сессия",
                f"{top_session.get('stationName')} · {top_session.get('date')} · {top_session.get('start')}",
                "text",
            )
        )
        analytics.append(("Сумма", _money_int(top_session.get("sum")), "money"))
    else:
        analytics.append(("Самая дорогая сессия", "—", "text"))
        analytics.append(("Сумма", 0, "money"))

    if longest_session:
        analytics.append(
            (
                "Самая длинная сессия",
                f"{longest_session.get('stationName')} · {longest_session.get('date')} · {longest_session.get('start')}",
                "text",
            )
        )
        analytics.append(
            (
                "Длительность",
                _minutes_label(_safe_int(longest_session.get("_durationMinutes"), 0)),
                "text",
            )
        )
    else:
        analytics.append(("Самая длинная сессия", "—", "text"))
        analytics.append(("Длительность", "—", "text"))

    analytics.append(("Самая прибыльная станция", top_station_name, "text"))
    analytics.append(("Выручка", top_station_revenue, "money"))
    return analytics
