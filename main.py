# -*- coding: utf-8 -*-
from __future__ import annotations

from flask import Flask, render_template, jsonify, request, send_file, url_for
import base64
import ctypes
import hashlib
import os
import socket
import sys
import threading
import time
import webbrowser
import json
import re
import shutil
from pathlib import Path
from datetime import datetime, timedelta
from io import BytesIO
import subprocess

try:
    import winreg
except Exception:  # pragma: no cover - non-Windows fallback for source runs
    winreg = None



APP_NAME = "PS Lounge"
MUTEX_NAME = r"Local\PSLounge_SingleInstance_v1"

URL_TXT = "PS_Lounge_URL.txt"
URL_SHORTCUT = "Open_PS_Lounge.url"
LOG_FILE = "PS_Lounge_Log.txt"
STATE_FILE = "pslounge_state.json"
STATE_BAK_FILE = "pslounge_state.bak.json"
STATE_HISTORY_DIR = "state_history"
LICENSE_FILE = "pslounge_license.json"
SCHEMA_VERSION = 3
LICENSE_SCHEMA_VERSION = 1
LICENSE_PRODUCT = "pslounge-desktop"
STATE_HISTORY_RETENTION_DAYS = 30
LICENSE_PUBLIC_MODULUS_B64 = "w6a7vadcU1zba6IcxF3TqJTOsKgP6HaHvn9cYjQt3miW8SdbPLMAcf3yckqQq5ZMy+NslOc/bS3XippfLVOVlzVIcE8iq4P42Y1M1/dgx3A9II9wsxAs5Ch6zXsTR+bifnUKvZtx7ZriA27N5p7AVtbXuHjxfejJ3sm4Qk7QymhwqjvYZ0tP1rDX0iXljTlyJIzGUO6/iguf324rTJ4dG9+zwHwnu2lCafDXa3MVyMPnwiEQ0W1XG+t69NuleLq0f3b15T82USI/5tGOfWlSiDWSpa4B9sV3JrFhsTV26v2veXSdwW/HBnn5W6DK/p9AtmBO9DcH3oSBbiRSs8TLtQ=="
LICENSE_PUBLIC_EXPONENT_B64 = "AQAB"

def _ru_plural(n: int, one: str, few: str, many: str) -> str:
    n = abs(int(n))
    if n % 10 == 1 and n % 100 != 11:
        return one
    if n % 10 in (2, 3, 4) and n % 100 not in (12, 13, 14):
        return few
    return many


def _minutes_label(total: int) -> str:
    total = max(0, int(total or 0))
    h = total // 60
    mi = total % 60
    return f"{h} ч {mi:02d} мин"


def _format_date_ru(value: str | int | float | datetime | None) -> str:
    if value is None:
        return ""
    try:
        if isinstance(value, datetime):
            d = value
        elif isinstance(value, (int, float)):
            d = datetime.fromtimestamp(float(value) / 1000 if float(value) > 10_000_000_000 else float(value))
        else:
            s = str(value).strip()
            if not s:
                return ""
            if re.fullmatch(r"\d{4}-\d{2}-\d{2}", s):
                d = datetime.strptime(s, "%Y-%m-%d")
            else:
                d = datetime.fromisoformat(s)
        return d.strftime("%d.%m.%Y")
    except Exception:
        return str(value)


def _hhmm(ms: int | float | None) -> str:
    if not ms:
        return ""
    try:
        dt = datetime.fromtimestamp(float(ms) / 1000)
        return dt.strftime("%H:%M")
    except Exception:
        return ""


def _money_int(v) -> int:
    try:
        return max(0, int(round(float(v or 0))))
    except Exception:
        return 0


def _money_label(v) -> str:
    return f"{_money_int(v):,} ₽".replace(",", " ")


def _payment_label(v) -> str:
    s = str(v or '').strip().lower()
    mapping = {
        'cash': 'Наличные',
        'card': 'Карта',
        'transfer': 'Перевод',
    }
    return mapping.get(s, 'Не указано')


def _station_type_label(v) -> str:
    s = str(v or '').strip().lower()
    mapping = {
        'ps': 'PlayStation',
        'playstation': 'PlayStation',
        'racing': 'Симулятор гонок',
        'simulator': 'Симулятор гонок',
        'switch': 'Nintendo Switch',
        'nintendo': 'Nintendo Switch',
    }
    return mapping.get(s, 'PlayStation')


def _station_name(sid: int | None, station_type: str | None = None) -> str:
    sid = int(sid or 0)
    st = str(station_type or '').strip().lower()
    if sid == 5 or st in {'racing', 'simulator'}:
        return 'Симулятор гонок'
    if sid == 6 or st in {'switch', 'nintendo'}:
        return 'Nintendo Switch'
    if 1 <= sid <= 4:
        return f'PS{sid}'
    return f'Станция {sid}' if sid else ''


def _tariff_label_export(rec: dict) -> str:
    label = str(rec.get('tariffLabel') or '').strip()
    if label:
        # Remove service suffixes like "(-1 час)" from export view.
        label = re.sub(r"\s*\([^)]*\)\s*$", "", label).strip()
        return label

    tid = str(rec.get('tariffId') or '').strip()
    if not tid:
        return ''

    m = re.match(r"^t(\d+)$", tid, flags=re.IGNORECASE)
    if m:
        n = int(m.group(1))
        return f"{n} {_ru_plural(n, 'час', 'часа', 'часов')}"

    m = re.match(r"^custom\s*[:\-]\s*(\d+)$", tid, flags=re.IGNORECASE)
    if m:
        total = int(m.group(1))
        h = total // 60
        mi = total % 60
        parts = []
        if h:
            parts.append(f"{h} {_ru_plural(h, 'час', 'часа', 'часов')}")
        if mi or not parts:
            parts.append(f"{mi} {_ru_plural(mi, 'минута', 'минуты', 'минут')}")
        return ' '.join(parts)

    return tid



def _safe_int(value, default: int = 0) -> int:
    try:
        return int(round(float(value)))
    except Exception:
        return default


def _normalize_payment_method(value, fallback: str = 'cash') -> str:
    raw = str(value or '').strip().lower()
    return raw if raw in {'cash', 'card', 'transfer'} else fallback


def _normalize_station_type_key(value, fallback: str = 'ps') -> str:
    raw = str(value or '').strip().lower()
    mapping = {
        'ps': 'ps',
        'playstation': 'ps',
        'simulator': 'simulator',
        'racing': 'simulator',
        'switch': 'switch',
        'nintendo': 'switch',
    }
    return mapping.get(raw, fallback)


def _parse_date_key(value: str | None, fallback: str | None = None) -> str:
    candidate = str(value or fallback or '').strip()
    try:
        return datetime.strptime(candidate, "%Y-%m-%d").strftime("%Y-%m-%d")
    except Exception as e:
        raise ValueError(f"invalid_date_key:{candidate}") from e


def _b64url_decode(value: str) -> bytes:
    raw = str(value or "").strip().encode("ascii")
    raw += b"=" * (-len(raw) % 4)
    return base64.urlsafe_b64decode(raw)


def _b64url_encode(value: bytes) -> str:
    return base64.urlsafe_b64encode(value).decode("ascii").rstrip("=")


def _registry_machine_guid() -> str:
    if winreg is None:
        return ""
    try:
        with winreg.OpenKey(winreg.HKEY_LOCAL_MACHINE, r"SOFTWARE\Microsoft\Cryptography") as key:
            value, _ = winreg.QueryValueEx(key, "MachineGuid")
            return str(value or "").strip()
    except Exception:
        return ""


def _system_drive_serial() -> str:
    drive = os.environ.get("SystemDrive", "C:").rstrip("\\/")
    try:
        completed = subprocess.run(
            ["cmd", "/c", "vol", drive],
            capture_output=True,
            text=True,
            timeout=5,
            check=False,
        )
        text = f"{completed.stdout}\n{completed.stderr}"
        match = re.search(r"([A-F0-9]{4}-[A-F0-9]{4})", text, flags=re.IGNORECASE)
        return match.group(1).upper() if match else ""
    except Exception:
        return ""


def _machine_fingerprint_parts() -> list[str]:
    parts = [
        _registry_machine_guid(),
        _system_drive_serial(),
        os.environ.get("PROCESSOR_IDENTIFIER", "").strip(),
        socket.gethostname().strip(),
    ]
    return [part for part in parts if part]


def _machine_fingerprint() -> str:
    joined = "|".join(_machine_fingerprint_parts())
    return hashlib.sha256(joined.encode("utf-8")).hexdigest()


def _machine_fingerprint_label() -> str:
    fp = _machine_fingerprint()
    return f"{fp[:8]}-{fp[8:16]}-{fp[16:24]}"


def _license_file_path() -> Path:
    return _app_dir() / LICENSE_FILE


def _license_public_key() -> tuple[int, int]:
    modulus = int.from_bytes(base64.b64decode(LICENSE_PUBLIC_MODULUS_B64), "big")
    exponent = int.from_bytes(base64.b64decode(LICENSE_PUBLIC_EXPONENT_B64), "big")
    return modulus, exponent


def _verify_license_token(token: str) -> dict:
    parts = str(token or "").strip().split(".")
    if len(parts) != 2:
        raise ValueError("invalid_token_format")
    payload_segment, signature_segment = parts
    payload_bytes = _b64url_decode(payload_segment)
    signature_bytes = _b64url_decode(signature_segment)
    payload = json.loads(payload_bytes.decode("utf-8"))
    if not isinstance(payload, dict):
        raise ValueError("invalid_payload")
    modulus, exponent = _license_public_key()
    sig_int = int.from_bytes(signature_bytes, "big")
    digest = hashlib.sha256(payload_segment.encode("ascii")).digest()
    expected = int.from_bytes(digest, "big")
    actual = pow(sig_int, exponent, modulus)
    if actual != expected:
        raise ValueError("bad_signature")
    if str(payload.get("product") or "").strip() != LICENSE_PRODUCT:
        raise ValueError("wrong_product")
    return payload


def _read_license_file() -> dict | None:
    path = _license_file_path()
    try:
        if not path.exists():
            return None
        data = json.loads(path.read_text(encoding="utf-8"))
        if not isinstance(data, dict):
            return None
        return data
    except Exception as e:
        log(f"Failed to read license file: {e!r}")
        return None


def _write_license_file(data: dict) -> None:
    path = _license_file_path()
    payload = json.dumps(data, ensure_ascii=False, indent=2)
    tmp = path.with_suffix(path.suffix + ".tmp")
    tmp.write_text(payload, encoding="utf-8")
    tmp.replace(path)


def _license_status() -> dict:
    machine_fp = _machine_fingerprint()
    base = {
        "ok": True,
        "licensed": False,
        "status": "missing",
        "machineFingerprint": machine_fp,
        "machineFingerprintLabel": _machine_fingerprint_label(),
    }
    record = _read_license_file()
    if not record:
        return base
    token = str(record.get("token") or "").strip()
    if not token:
        return {**base, "status": "corrupt", "detail": "empty_token"}
    try:
        payload = _verify_license_token(token)
    except ValueError as e:
        return {**base, "status": "invalid", "detail": str(e)}
    expected_fp = str(payload.get("fingerprint") or "").strip().lower()
    if not expected_fp or expected_fp != machine_fp.lower():
        return {**base, "status": "device_mismatch", "detail": "fingerprint_mismatch"}
    customer = str(payload.get("customer") or "").strip()
    return {
        **base,
        "licensed": True,
        "status": "active",
        "customer": customer,
        "issuedAt": payload.get("issuedAt"),
        "license": {
            "customer": customer,
            "product": payload.get("product"),
            "fingerprint": expected_fp,
        },
    }


def _activate_license_token(token: str) -> dict:
    payload = _verify_license_token(token)
    machine_fp = _machine_fingerprint()
    expected_fp = str(payload.get("fingerprint") or "").strip().lower()
    if not expected_fp:
        raise ValueError("missing_fingerprint")
    if expected_fp != machine_fp.lower():
        raise ValueError("fingerprint_mismatch")
    record = {
        "schemaVersion": LICENSE_SCHEMA_VERSION,
        "product": LICENSE_PRODUCT,
        "token": token.strip(),
        "activatedAt": datetime.now().isoformat(timespec="seconds"),
    }
    _write_license_file(record)
    return _license_status()


def _require_license():
    status = _license_status()
    if status.get("licensed"):
        return None
    return jsonify({
        "ok": False,
        "error": "license_required",
        "license": {
            "status": status.get("status"),
            "detail": status.get("detail", ""),
            "machineFingerprint": status.get("machineFingerprint"),
            "machineFingerprintLabel": status.get("machineFingerprintLabel"),
        },
    }), 403


def _normalize_station_definition(defn: object, idx: int) -> dict | None:
    if not isinstance(defn, dict):
        return None
    raw_id = _safe_int(defn.get("id"), idx + 1)
    station_type = _normalize_station_type_key(defn.get("type"), "ps")
    name = str(defn.get("name") or _station_name(raw_id, station_type)).strip()
    if raw_id <= 0 or not name:
        return None
    return {"id": raw_id, "name": name, "type": station_type}


def _state_station_definitions(state: dict | None) -> list[dict]:
    settings = state.get("settings") if isinstance(state, dict) else None
    raw_defs = settings.get("stationDefinitions") if isinstance(settings, dict) else None
    if isinstance(raw_defs, list):
        defs = []
        for idx, item in enumerate(raw_defs):
            normalized = _normalize_station_definition(item, idx)
            if normalized is not None:
                defs.append(normalized)
        if defs:
            defs.sort(key=lambda item: item["id"])
            return defs
    stations = state.get("stations") if isinstance(state, dict) else None
    defs = []
    if isinstance(stations, list):
        for idx, item in enumerate(stations):
            if not isinstance(item, dict):
                continue
            sid = _safe_int(item.get("id"), idx + 1)
            stype = _normalize_station_type_key(item.get("stationType"), "simulator" if sid == 5 else "switch" if sid == 6 else "ps")
            defs.append({
                "id": sid,
                "name": str(item.get("name") or _station_name(sid, stype)).strip(),
                "type": stype,
            })
    return sorted(defs, key=lambda item: item["id"])


def _report_station_names(state: dict | None, session_rows: list[dict], sales_rows: list[dict]) -> list[str]:
    names = []
    seen = set()
    for item in _state_station_definitions(state):
        name = str(item.get("name") or "").strip()
        if name and name not in seen:
            names.append(name)
            seen.add(name)
    for row in [*session_rows, *sales_rows]:
        name = str(row.get("stationName") or "").strip()
        if name and name not in seen:
            names.append(name)
            seen.add(name)
    return names


def _is_valid_sale_item(value: object) -> bool:
    if not isinstance(value, dict):
        return False
    if value.get("minutes") is not None and not isinstance(value.get("minutes"), (int, float)):
        return False
    if value.get("amount") is not None and not isinstance(value.get("amount"), (int, float)):
        return False
    if value.get("time") is not None and not isinstance(value.get("time"), (int, float)):
        return False
    return True


def _is_valid_session_record(value: object) -> bool:
    if not isinstance(value, dict):
        return False
    if value.get("stationId") is not None and not isinstance(value.get("stationId"), (int, float)):
        return False
    if value.get("startTime") is not None and not isinstance(value.get("startTime"), (int, float)):
        return False
    if value.get("endTime") is not None and not isinstance(value.get("endTime"), (int, float)):
        return False
    sales = value.get("sales")
    if sales is not None:
        if not isinstance(sales, list):
            return False
        for sale in sales:
            if not _is_valid_sale_item(sale):
                return False
    return True


def _compute_session_duration_minutes(rec: dict, paid_minutes_total: int, grace_minutes: int, overdue_minutes: int) -> tuple[int | None, bool]:
    start_ts = _safe_int(rec.get('startTime'), 0)
    end_ts = _safe_int(rec.get('endTime'), 0)
    if start_ts <= 0 or end_ts <= 0:
        return None, False

    adjusted = False
    if end_ts < start_ts:
        # Defensive correction for sessions that crossed midnight but were persisted on the same logical day.
        if (start_ts - end_ts) <= 24 * 60 * 60 * 1000:
            end_ts += 24 * 60 * 60 * 1000
            adjusted = True
        else:
            return 0, True

    elapsed_minutes = max(0, round((end_ts - start_ts) / 60000))
    if elapsed_minutes <= 0:
        return 0, adjusted

    paid_minutes_total = max(0, _safe_int(paid_minutes_total, 0))
    allowed_overrun = max(0, _safe_int(grace_minutes, 0)) + max(0, _safe_int(overdue_minutes, 0))

    if paid_minutes_total > 0:
        hard_cap = max(paid_minutes_total + allowed_overrun + 15, round(paid_minutes_total * 1.5))
        if elapsed_minutes > hard_cap:
            return paid_minutes_total, True

    return elapsed_minutes, adjusted


def _build_session_sales_rows(rec: dict, key: str, sid: int, stype: str, display_station: str, display_type: str, session_key: str) -> tuple[list[dict], int, int, str, str, list[str]]:
    sales_rows: list[dict] = []
    payment_values: list[str] = []
    paid_minutes_total = 0
    fallback_payment = _normalize_payment_method(rec.get('paymentMethod'), 'cash')
    fallback_ts = _safe_int(rec.get('startTime'), 0)

    sales = rec.get('sales') if isinstance(rec.get('sales'), list) else []
    for sale in sales:
        if not isinstance(sale, dict):
            continue

        sale_minutes = max(0, _safe_int(sale.get('minutes'), 0))
        sale_amount = _money_int(sale.get('amount'))
        sale_ts = _safe_int(sale.get('time'), fallback_ts) or fallback_ts
        sale_payment = _normalize_payment_method(sale.get('paymentMethod'), fallback_payment)
        operation = str(sale.get('label') or sale.get('type') or 'Начисление').strip() or 'Начисление'

        if sale_amount <= 0 and sale_minutes <= 0 and operation == 'Начисление':
            continue

        paid_minutes_total += sale_minutes
        payment_values.append(sale_payment)

        sales_rows.append({
            'date': _format_date_ru(datetime.fromtimestamp(sale_ts / 1000).strftime('%Y-%m-%d')) if sale_ts > 0 else _format_date_ru(key),
            'time': _hhmm(sale_ts) if sale_ts > 0 else '—',
            'stationId': sid,
            'stationName': display_station,
            'stationType': display_type,
            'operation': operation,
            'minutes': sale_minutes,
            'amount': sale_amount,
            'payment': _payment_label(sale_payment),
            '_stationTypeRaw': stype,
            '_paymentRaw': sale_payment,
            '_sessionKey': session_key,
        })

    if sales_rows:
        total_amount = sum(row['amount'] for row in sales_rows)
    else:
        total_amount = _money_int(rec.get('totalAmount'))
        if total_amount > 0:
            sales_rows.append({
                'date': _format_date_ru(key),
                'time': _hhmm(fallback_ts) if fallback_ts > 0 else '—',
                'stationId': sid,
                'stationName': display_station,
                'stationType': display_type,
                'operation': 'Продажа',
                'minutes': 0,
                'amount': total_amount,
                'payment': _payment_label(fallback_payment),
                '_stationTypeRaw': stype,
                '_paymentRaw': fallback_payment,
                '_sessionKey': session_key,
            })
            payment_values.append(fallback_payment)

    integrity_flags: list[str] = []
    raw_total_amount = _money_int(rec.get('totalAmount'))
    if raw_total_amount > 0 and sales_rows and raw_total_amount != total_amount:
        integrity_flags.append('sales_total_mismatch')

    if payment_values:
        uniq = []
        for pm in payment_values:
            if pm not in uniq:
                uniq.append(pm)
        payment_raw = uniq[0] if len(uniq) == 1 else 'mixed'
        payment_method = _payment_label(uniq[0]) if len(uniq) == 1 else 'Смешанная'
    else:
        payment_raw = fallback_payment
        payment_method = _payment_label(payment_raw)

    return sales_rows, total_amount, paid_minutes_total, payment_raw, payment_method, integrity_flags


def _collect_report_rows(state: dict, from_key: str, to_key: str, station_type: str = '', payment: str = '') -> tuple[list[dict], list[dict], list[str]]:
    sessions_map = state.get('sessions') or {}
    settings = state.get('settings') if isinstance(state.get('settings'), dict) else {}
    grace_minutes = max(0, _safe_int(settings.get('graceMinutes'), 0))
    overdue_minutes = max(0, _safe_int(settings.get('overdueMinutes'), 0))
    keys = sorted(k for k in sessions_map.keys() if from_key <= k <= to_key)
    session_rows: list[dict] = []
    sales_rows: list[dict] = []
    station_type = _normalize_station_type_key(station_type, '')
    payment = str(payment or '').strip().lower()

    for key in keys:
        for rec_index, rec in enumerate(sessions_map.get(key) or []):
            if not isinstance(rec, dict):
                continue

            sid = _safe_int(rec.get('stationId'), 0)
            stype = _normalize_station_type_key(rec.get('stationType') or ('racing' if sid == 5 else 'switch' if sid == 6 else 'ps'))
            display_type = _station_type_label(stype)
            display_station = rec.get('stationName') or _station_name(sid, stype)

            session_key = f"{key}|{sid}|{_safe_int(rec.get('startTime'), 0)}|{_safe_int(rec.get('endTime'), 0)}|{rec_index}"
            session_sales_rows, total_amount, paid_minutes_total, payment_raw, payment_method, integrity_flags = _build_session_sales_rows(
                rec, key, sid, stype, display_station, display_type, session_key
            )
            sales_rows.extend(session_sales_rows)

            duration_minutes, duration_adjusted = _compute_session_duration_minutes(
                rec,
                paid_minutes_total,
                grace_minutes,
                overdue_minutes,
            )

            row = {
                'date': _format_date_ru(key),
                'stationId': sid,
                'stationName': display_station,
                'stationType': display_type,
                'paidTime': _minutes_label(paid_minutes_total),
                'start': _hhmm(rec.get('startTime')),
                'finish': _hhmm(rec.get('endTime')) if rec.get('endTime') else '—',
                'duration': _minutes_label(duration_minutes) if duration_minutes is not None else 'Активна',
                'sum': total_amount,
                'payment': payment_method,
                'close': 'идёт' if not rec.get('endTime') else ('авто' if rec.get('mode') == 'auto' else 'вручную'),
                '_stationTypeRaw': stype,
                '_paymentRaw': payment_raw,
                '_paidMinutes': paid_minutes_total,
                '_endTime': rec.get('endTime'),
                '_startTime': rec.get('startTime'),
                '_durationMinutes': duration_minutes or 0,
                '_durationAdjusted': duration_adjusted,
                '_integrityFlags': integrity_flags + (['duration_adjusted'] if duration_adjusted else []),
                '_sessionKey': session_key,
            }
            session_rows.append(row)

    if station_type and station_type != 'all':
        session_rows = [r for r in session_rows if r['_stationTypeRaw'] == station_type]
        sales_rows = [r for r in sales_rows if r['_stationTypeRaw'] == station_type]
    if payment and payment != 'all':
        session_rows = [r for r in session_rows if r['_paymentRaw'] == payment or (payment == 'mixed' and r['payment'] == 'Смешанная')]
        sales_rows = [r for r in sales_rows if r['_paymentRaw'] == payment]

    return session_rows, sales_rows, keys


def _excel_styles():
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter

    header_fill = PatternFill('solid', fgColor='1F2A44')
    subheader_fill = PatternFill('solid', fgColor='EAF0FF')
    section_fill = PatternFill('solid', fgColor='F3F6FB')
    total_fill = PatternFill('solid', fgColor='EEF2F8')
    header_font = Font(bold=True, color='FFFFFF')
    title_font = Font(bold=True, size=16)
    subtitle_font = Font(bold=True, size=12)
    bold_font = Font(bold=True)
    thin = Side(style='thin', color='D9DEE8')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal='left', vertical='center')
    left = Alignment(horizontal='left', vertical='center')

    def style_header(ws, row_idx):
        for cell in ws[row_idx]:
            if cell.value in (None, ''):
                continue
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center
            cell.border = border

    def style_range(ws, min_row, max_row, min_col, max_col, fill=None, font=None, alignment=None):
        for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
            for cell in row:
                if fill:
                    cell.fill = fill
                if font:
                    cell.font = font
                if alignment:
                    cell.alignment = alignment
                cell.border = border

    def auto_width(ws, min_width=10, max_width=32):
        dims = {}
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                val = str(cell.value)
                dims[cell.column] = max(dims.get(cell.column, 0), len(val))
                if cell.border is None:
                    cell.border = border
        for idx, length in dims.items():
            ws.column_dimensions[get_column_letter(idx)].width = min(max(length + 2, min_width), max_width)

    return {
        'header_fill': header_fill,
        'subheader_fill': subheader_fill,
        'section_fill': section_fill,
        'total_fill': total_fill,
        'header_font': header_font,
        'title_font': title_font,
        'subtitle_font': subtitle_font,
        'bold_font': bold_font,
        'border': border,
        'center': center,
        'left': left,
        'style_header': style_header,
        'style_range': style_range,
        'auto_width': auto_width,
    }


EXCEL_RUB_FORMAT = '# ##0" ₽"'
EXCEL_INT_FORMAT = '0'


def _apply_number_format_range(ws, col_idx: int, start_row: int, end_row: int, fmt: str):
    if end_row < start_row:
        return
    for row in range(start_row, end_row + 1):
        cell = ws.cell(row=row, column=col_idx)
        if cell.value in (None, ''):
            continue
        if isinstance(cell.value, (int, float)):
            cell.number_format = fmt


def _add_excel_table(ws, table_name: str, start_row: int, data_end_row: int, end_col: int, totals: dict[int, str] | None = None):
    from openpyxl.worksheet.table import Table, TableStyleInfo, TableColumn
    from openpyxl.utils import get_column_letter

    table_end_row = data_end_row + (1 if totals else 0)
    ref = f"A{start_row}:{get_column_letter(end_col)}{table_end_row}"
    tab = Table(displayName=table_name, ref=ref)
    tab.tableStyleInfo = TableStyleInfo(
        name='TableStyleMedium2',
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )

    headers = [ws.cell(start_row, col_idx).value or f'Column{col_idx}' for col_idx in range(1, end_col + 1)]
    tab.tableColumns = [TableColumn(id=idx, name=str(name)) for idx, name in enumerate(headers, start=1)]
    if totals:
        tab.totalsRowCount = 1
        tab.totalsRowShown = True
        for col_idx, func in totals.items():
            if 1 <= col_idx <= len(tab.tableColumns):
                tab.tableColumns[col_idx - 1].totalsRowFunction = func
    ws.add_table(tab)
    return table_end_row


def _write_data_sheet(
    ws,
    headers: list[str],
    rows: list[list],
    table_name: str,
    money_cols: tuple[int, ...] = (),
    int_cols: tuple[int, ...] = (),
    totals: dict[int, str] | None = None,
):
    S = _excel_styles()

    ws.append(headers)
    S['style_header'](ws, 1)

    for row in rows:
        ws.append(row)

    data_end_row = ws.max_row
    table_end_row = data_end_row
    if rows:
        table_end_row = _add_excel_table(ws, table_name, 1, data_end_row, len(headers), totals=totals)

    start_fmt_row = 2
    end_fmt_row = data_end_row

    for col_idx in money_cols:
        _apply_number_format_range(ws, col_idx, start_fmt_row, end_fmt_row, EXCEL_RUB_FORMAT)
        if totals and rows:
            ws.cell(row=table_end_row, column=col_idx).number_format = EXCEL_RUB_FORMAT

    for col_idx in int_cols:
        _apply_number_format_range(ws, col_idx, start_fmt_row, end_fmt_row, EXCEL_INT_FORMAT)
        if totals and rows:
            ws.cell(row=table_end_row, column=col_idx).number_format = EXCEL_INT_FORMAT

    S['auto_width'](ws)
    from openpyxl.styles import Alignment
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=len(headers)):
        for cell in row:
            cell.alignment = Alignment(horizontal='left', vertical='center')
    ws.freeze_panes = 'A2'



def _add_revenue_chart(ws, min_row: int, max_row: int, anchor: str = 'D12'):
    from openpyxl.chart import BarChart, Reference
    from openpyxl.chart.label import DataLabelList

    if max_row <= min_row:
        return

    chart = BarChart()
    chart.type = 'col'
    chart.style = 11
    chart.title = 'Выручка по дням'
    chart.height = 7.2
    chart.width = 12.8
    chart.legend = None
    chart.gapWidth = 65
    chart.overlap = 0
    chart.varyColors = False

    chart.y_axis.title = 'Выручка, ₽'
    chart.y_axis.number_format = '# ##0'
    try:
        chart.y_axis.majorGridlines = None
    except Exception:
        pass
    try:
        chart.x_axis.majorGridlines = None
    except Exception:
        pass
    chart.x_axis.tickLblPos = 'low'
    chart.x_axis.delete = False
    try:
        chart.x_axis.title = None
    except Exception:
        pass

    data = Reference(ws, min_col=2, min_row=min_row, max_row=max_row)
    cats = Reference(ws, min_col=1, min_row=min_row + 1, max_row=max_row)

    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)

    if chart.series:
        series = chart.series[0]
        series.graphicalProperties.line.noFill = True
        try:
            series.graphicalProperties.solidFill = '4F81BD'
        except Exception:
            pass

    chart.dLbls = DataLabelList()
    chart.dLbls.showVal = True
    chart.dLbls.showCatName = False
    chart.dLbls.showLegendKey = False
    chart.dLbls.showSerName = False
    chart.dLbls.showPercent = False
    chart.dLbls.position = 'outEnd'

    ws.add_chart(chart, anchor)


def _set_summary_value(cell, value, kind: str):
    cell.value = value
    if kind == 'money':
        cell.number_format = EXCEL_RUB_FORMAT
    elif kind == 'int':
        cell.number_format = EXCEL_INT_FORMAT


def _apply_daily_heatmap(ws, col_idx: int, start_row: int, end_row: int):
    from openpyxl.styles import PatternFill, Font
    values = []
    for row in range(start_row, end_row + 1):
        val = ws.cell(row=row, column=col_idx).value
        if isinstance(val, (int, float)):
            values.append(float(val))
    if not values:
        return
    vmax = max(values) or 1
    for row in range(start_row, end_row + 1):
        cell = ws.cell(row=row, column=col_idx)
        val = cell.value if isinstance(cell.value, (int, float)) else 0
        ratio = max(0.0, min(1.0, float(val) / vmax if vmax else 0.0))
        if ratio >= 0.85:
            fill = PatternFill('solid', fgColor='C6EFCE')
        elif ratio >= 0.55:
            fill = PatternFill('solid', fgColor='DDEBF7')
        elif ratio > 0:
            fill = PatternFill('solid', fgColor='FFF2CC')
        else:
            fill = PatternFill('solid', fgColor='F3F6FB')
        cell.fill = fill
        if ratio >= 0.85:
            cell.font = Font(bold=True)

def _add_station_load_chart(ws, min_row: int, max_row: int, anchor: str = 'G12'):
    from openpyxl.chart import BarChart, Reference
    from openpyxl.chart.label import DataLabelList

    if max_row <= min_row:
        return

    chart = BarChart()
    chart.type = 'bar'
    chart.style = 10
    chart.title = 'Загрузка станций'
    chart.height = 6.4
    chart.width = 8.6
    chart.legend = None
    chart.gapWidth = 40
    chart.varyColors = False

    try:
        chart.y_axis.majorGridlines = None
    except Exception:
        pass
    try:
        chart.x_axis.majorGridlines = None
    except Exception:
        pass
    chart.x_axis.title = 'Часы'
    chart.x_axis.number_format = '0.0'
    chart.y_axis.title = None

    data = Reference(ws, min_col=8, min_row=min_row, max_row=max_row)
    cats = Reference(ws, min_col=7, min_row=min_row + 1, max_row=max_row)

    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    if chart.series:
        series = chart.series[0]
        series.graphicalProperties.line.noFill = True
        try:
            series.graphicalProperties.solidFill = '6FA8DC'
        except Exception:
            pass

    chart.dLbls = DataLabelList()
    chart.dLbls.showVal = True
    chart.dLbls.position = 'outEnd'
    ws.add_chart(chart, anchor)


def _build_shift_workbook(day_key: str, session_rows: list[dict], sales_rows: list[dict], state: dict | None = None):
    from openpyxl import Workbook

    S = _excel_styles()
    wb = Workbook()
    ws_shift = wb.active
    ws_shift.title = 'Смена'
    ws_sessions = wb.create_sheet('Сессии за день')
    ws_sales = wb.create_sheet('Начисления за день')

    # Sheet 1: daily operational shift summary
    ws_shift.merge_cells('A1:D1')
    ws_shift['A1'] = 'PS Lounge — Смена'
    ws_shift['A1'].font = S['title_font']
    ws_shift['A3'] = 'Дата'
    ws_shift['B3'] = _format_date_ru(day_key)
    ws_shift['A3'].font = S['bold_font']
    ws_shift['A5'] = 'Показатель'
    ws_shift['B5'] = 'Значение'
    S['style_header'](ws_shift, 5)

    closed_sessions = [r for r in session_rows if r.get('_endTime')]
    revenue = sum(r['amount'] for r in sales_rows)
    session_count = len(session_rows)
    avg_check = round(revenue / session_count) if session_count else 0
    avg_minutes = round(sum(max(0, int(r.get('_durationMinutes') or 0)) for r in closed_sessions) / len(closed_sessions)) if closed_sessions else 0
    summary_rows = [
        ('Выручка за день', revenue, 'money'),
        ('Сессий', session_count, 'int'),
        ('Средний чек', avg_check, 'money'),
        ('Средняя длительность', _minutes_label(avg_minutes), 'text'),
    ]
    row = 6
    for label, val, kind in summary_rows:
        ws_shift[f'A{row}'] = label
        ws_shift[f'A{row}'].font = S['bold_font']
        _set_summary_value(ws_shift[f'B{row}'], val, kind)
        row += 1

    row += 1
    ws_shift[f'A{row}'] = 'Выручка по станциям'
    ws_shift[f'A{row}'].font = S['subtitle_font']
    row += 1
    ws_shift[f'A{row}'] = 'Станция'
    ws_shift[f'B{row}'] = 'Выручка'
    S['style_header'](ws_shift, row)
    row += 1

    station_order = [
        ('PS1', 'PS1'), ('PS2', 'PS2'), ('PS3', 'PS3'), ('PS4', 'PS4'),
        ('Симулятор гонок', 'Симулятор гонок'), ('Nintendo Switch', 'Nintendo Switch')
    ]
    station_totals = {name: 0 for _, name in station_order}
    for sale in sales_rows:
        station_totals[sale['stationName']] = station_totals.get(sale['stationName'], 0) + sale['amount']
    for _, name in station_order:
        ws_shift[f'A{row}'] = name
        _set_summary_value(ws_shift[f'B{row}'], station_totals.get(name, 0), 'money')
        row += 1

    row += 1
    ws_shift[f'A{row}'] = 'Выручка по оплате'
    ws_shift[f'A{row}'].font = S['subtitle_font']
    row += 1
    ws_shift[f'A{row}'] = 'Оплата'
    ws_shift[f'B{row}'] = 'Сумма'
    S['style_header'](ws_shift, row)
    row += 1
    payment_totals = {'Наличные': 0, 'Карта': 0, 'Перевод': 0}
    for sale in sales_rows:
        payment_totals[sale['payment']] = payment_totals.get(sale['payment'], 0) + sale['amount']
    for label in ('Наличные', 'Карта', 'Перевод'):
        ws_shift[f'A{row}'] = label
        _set_summary_value(ws_shift[f'B{row}'], payment_totals.get(label, 0), 'money')
        row += 1

    row += 1
    ws_shift[f'A{row}'] = 'Активные сессии сейчас'
    ws_shift[f'A{row}'].font = S['subtitle_font']
    row += 1
    ws_shift[f'A{row}'] = 'Станция'
    ws_shift[f'B{row}'] = 'Оплачено времени'
    ws_shift[f'C{row}'] = 'Окончание'
    ws_shift[f'D{row}'] = 'Оплачено'
    S['style_header'](ws_shift, row)
    row += 1
    active = [r for r in session_rows if not r.get('_endTime')]
    if active:
        for r in active:
            ws_shift[f'A{row}'] = r['stationName']
            ws_shift[f'B{row}'] = r['paidTime']
            ws_shift[f'C{row}'] = '—'
            _set_summary_value(ws_shift[f'D{row}'], r['sum'], 'money')
            row += 1
    else:
        ws_shift[f'A{row}'] = 'Нет активных сессий'
        row += 1

    S['style_range'](ws_shift, 3, ws_shift.max_row, 1, 4, alignment=S['left'])
    S['style_header'](ws_shift, 5)
    S['auto_width'](ws_shift)
    ws_shift.freeze_panes = 'A6'

    # Sheet 2: sessions for day
    sess_headers = ['Дата', 'Станция', 'Тип станции', 'Оплачено времени', 'Старт', 'Финиш', 'Длительность', 'Сумма', 'Оплата', 'Завершение']
    session_table_rows = [
        [r['date'], r['stationName'], r['stationType'], r['paidTime'], r['start'], r['finish'], r['duration'], r['sum'], r['payment'], r['close']]
        for r in session_rows
    ]
    _write_data_sheet(
        ws_sessions,
        headers=sess_headers,
        rows=session_table_rows,
        table_name='ShiftSessionsTable',
        money_cols=(8,),
        totals={8: 'sum'},
    )

    # Sheet 3: sales for day
    sale_headers = ['Дата', 'Время', 'Станция', 'Тип станции', 'Операция', 'Минуты', 'Сумма', 'Оплата']
    sales_table_rows = [
        [r['date'], r['time'], r['stationName'], r['stationType'], r['operation'], r['minutes'], r['amount'], r['payment']]
        for r in sales_rows
    ]
    _write_data_sheet(
        ws_sales,
        headers=sale_headers,
        rows=sales_table_rows,
        table_name='ShiftSalesTable',
        money_cols=(7,),
        int_cols=(6,),
        totals={7: 'sum'},
    )

    return wb





def _build_audit_rows(session_rows: list[dict], sales_rows: list[dict]) -> tuple[list[list], int, int]:
    sales_by_session: dict[str, dict[str, int]] = {}
    orphan_sales = 0

    for sale in sales_rows:
        session_key = str(sale.get('_sessionKey') or '').strip()
        if not session_key:
            orphan_sales += 1
            continue
        bucket = sales_by_session.setdefault(session_key, {'amount': 0, 'minutes': 0, 'count': 0})
        bucket['amount'] += _money_int(sale.get('amount'))
        bucket['minutes'] += max(0, _safe_int(sale.get('minutes'), 0))
        bucket['count'] += 1

    rows: list[list] = []
    error_count = 0
    warning_count = 0

    for row in session_rows:
        session_key = str(row.get('_sessionKey') or '')
        sale_info = sales_by_session.get(session_key, {'amount': 0, 'minutes': 0, 'count': 0})
        session_sum = _money_int(row.get('sum'))
        sales_sum = int(sale_info['amount'])
        paid_minutes = max(0, _safe_int(row.get('_paidMinutes'), 0))
        sales_minutes = int(sale_info['minutes'])
        duration_minutes = max(0, _safe_int(row.get('_durationMinutes'), 0))
        close_mode = str(row.get('close') or '').strip().lower()
        flags = list(dict.fromkeys(row.get('_integrityFlags') or []))

        findings: list[tuple[str, str]] = []

        if sales_sum > 0 and session_sum != sales_sum:
            findings.append(('Ошибка', 'Сумма сессии не совпадает с суммой начислений'))
        if sales_minutes > 0 and paid_minutes > 0 and paid_minutes != sales_minutes:
            findings.append(('Ошибка', 'Оплаченные минуты не совпадают с минутами начислений'))
        if session_sum < 0 or sales_sum < 0:
            findings.append(('Ошибка', 'Обнаружена отрицательная сумма'))
        if paid_minutes < 0 or sales_minutes < 0:
            findings.append(('Ошибка', 'Обнаружено отрицательное количество минут'))

        if row.get('_endTime') and session_sum > 0 and duration_minutes == 0:
            findings.append(('Предупреждение', 'Нулевая длительность при ненулевой сумме'))
        if row.get('_endTime') and row.get('_startTime') and close_mode == 'вручную' and 0 <= duration_minutes <= 1 and session_sum > 0:
            findings.append(('Предупреждение', 'Ручное закрытие почти сразу после старта'))
        if row.get('_endTime') and paid_minutes > 0 and duration_minutes > paid_minutes + 60:
            findings.append(('Предупреждение', 'Фактическая длительность заметно больше оплаченного времени'))

        if 'sales_total_mismatch' in flags and not any(desc == 'Сумма сессии не совпадает с суммой начислений' for _, desc in findings):
            findings.append(('Ошибка', 'Итог сессии не совпадает с raw totalAmount'))
        if 'duration_adjusted' in flags and row.get('_endTime'):
            findings.append(('Предупреждение', 'Длительность была нормализована при экспорте'))

        seen: set[tuple[str, str]] = set()
        for level, description in findings:
            if (level, description) in seen:
                continue
            seen.add((level, description))
            if level == 'Ошибка':
                error_count += 1
            else:
                warning_count += 1
            rows.append([
                level,
                row.get('date'),
                row.get('stationName'),
                row.get('stationType'),
                row.get('payment'),
                session_sum,
                sales_sum,
                paid_minutes,
                sales_minutes,
                row.get('duration'),
                description,
            ])

    if orphan_sales:
        error_count += 1
        rows.append([
            'Ошибка',
            '',
            '—',
            '—',
            '—',
            0,
            0,
            0,
            0,
            '—',
            f'Найдены начисления без привязки к сессии: {orphan_sales}',
        ])

    return rows, error_count, warning_count


def _write_audit_sheet(wb, session_rows: list[dict], sales_rows: list[dict]) -> tuple[int, int]:
    from openpyxl.styles import Alignment, PatternFill, Font
    rows, error_count, warning_count = _build_audit_rows(session_rows, sales_rows)
    if not rows:
        return 0, 0

    ws = wb.create_sheet('Audit')
    headers = ['Уровень', 'Дата', 'Станция', 'Тип станции', 'Оплата', 'Сумма сессии', 'Сумма начислений', 'Оплачено минут', 'Минут начислено', 'Длительность', 'Описание']
    _write_data_sheet(ws, headers=headers, rows=rows, table_name='AuditTable', money_cols=(6, 7), int_cols=(8, 9))
    ws.column_dimensions['A'].width = 20
    ws.column_dimensions['K'].width = 55

    warn_fill = PatternFill('solid', fgColor='FFD966')
    err_fill = PatternFill('solid', fgColor='F4B183')
    bold_font = Font(bold=True)

    for row in range(2, ws.max_row + 1):
        level_cell = ws.cell(row=row, column=1)
        text = str(level_cell.value or '').strip()
        if text == 'Предупреждение':
            level_cell.fill = warn_fill
        elif text == 'Ошибка':
            level_cell.fill = err_fill
        level_cell.font = bold_font
        for col in range(1, 12):
            ws.cell(row=row, column=col).alignment = Alignment(horizontal='left', vertical='center')
        ws.cell(row=row, column=11).alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    return error_count, warning_count

def _top_analytics(session_rows: list[dict], sales_rows: list[dict]) -> list[tuple[str, object, str]]:
    top_session = max(session_rows, key=lambda r: (_money_int(r.get('sum')), _safe_int(r.get('_paidMinutes'), 0), str(r.get('stationName') or '')), default=None)
    closed_sessions = [r for r in session_rows if r.get('_endTime')]
    longest_session = max(closed_sessions, key=lambda r: (_safe_int(r.get('_durationMinutes'), 0), _money_int(r.get('sum')), str(r.get('stationName') or '')), default=None)

    station_revenue: dict[str, int] = {}
    for sale in sales_rows:
        name = str(sale.get('stationName') or '—')
        station_revenue[name] = station_revenue.get(name, 0) + _money_int(sale.get('amount'))
    top_station_name = '—'
    top_station_revenue = 0
    if station_revenue:
        top_station_name, top_station_revenue = max(station_revenue.items(), key=lambda item: (item[1], item[0]))

    analytics: list[tuple[str, object, str]] = []
    if top_session:
        analytics.append(('Самая дорогая сессия', f"{top_session.get('stationName')} · {top_session.get('date')} · {top_session.get('start')}", 'text'))
        analytics.append(('Сумма', _money_int(top_session.get('sum')), 'money'))
    else:
        analytics.append(('Самая дорогая сессия', '—', 'text'))
        analytics.append(('Сумма', 0, 'money'))

    if longest_session:
        analytics.append(('Самая длинная сессия', f"{longest_session.get('stationName')} · {longest_session.get('date')} · {longest_session.get('start')}", 'text'))
        analytics.append(('Длительность', _minutes_label(_safe_int(longest_session.get('_durationMinutes'), 0)), 'text'))
    else:
        analytics.append(('Самая длинная сессия', '—', 'text'))
        analytics.append(('Длительность', '—', 'text'))

    analytics.append(('Самая прибыльная станция', top_station_name, 'text'))
    analytics.append(('Выручка', top_station_revenue, 'money'))
    return analytics



def _build_report_workbook(from_key: str, to_key: str, session_rows: list[dict], sales_rows: list[dict], period_keys: list[str]):
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill, Font, Alignment
    from openpyxl.formatting.rule import DataBarRule

    S = _excel_styles()
    wb = Workbook()
    ws_summary = wb.active
    ws_summary.title = 'Сводка'
    ws_sessions = wb.create_sheet('Сессии')
    ws_sales = wb.create_sheet('Начисления')

    # Summary layout
    ws_summary.merge_cells('A1:B1')
    ws_summary['A1'] = 'PS Lounge'
    ws_summary['A1'].font = Font(bold=True, size=18)
    ws_summary['A1'].alignment = Alignment(horizontal='left', vertical='center')

    ws_summary.merge_cells('A2:B2')
    ws_summary['A2'] = 'Отчёт за период'
    ws_summary['A2'].font = Font(bold=True, size=12, color='666666')
    ws_summary['A2'].alignment = Alignment(horizontal='left', vertical='center')

    ws_summary['A4'] = 'Период'
    ws_summary['B4'] = _format_date_ru(from_key) if from_key == to_key else f"{_format_date_ru(from_key)} — {_format_date_ru(to_key)}"
    ws_summary['A4'].font = S['bold_font']
    # Проверка данных переносим вниз, чтобы не перегружать верхнюю часть

    closed_sessions = [r for r in session_rows if r.get('_endTime')]
    revenue = sum(r['amount'] for r in sales_rows)
    session_count = len(session_rows)
    avg_check = round(revenue / session_count) if session_count else 0
    avg_minutes = round(sum(max(0, int(r.get('_durationMinutes') or 0)) for r in closed_sessions) / len(closed_sessions)) if closed_sessions else 0

    # KPI cards
    kpi_fill = PatternFill('solid', fgColor='F7FAFF')
    kpi_value_font = Font(bold=True, size=14)
    kpi_title_font = Font(bold=True, size=11, color='44546A')
    kpis = [
        ('Выручка', revenue, 'money'),
        ('Сессий', session_count, 'int'),
        ('Средний чек', avg_check, 'money'),
        ('Средняя длительность', _minutes_label(avg_minutes), 'text'),
    ]
    kpi_rows = [(6,1,2),(6,3,4),(9,1,2),(9,3,4)]
    for (label, val, kind), (r, c1, c2) in zip(kpis, kpi_rows):
        ws_summary.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
        ws_summary.merge_cells(start_row=r+1, start_column=c1, end_row=r+1, end_column=c2)
        title_cell = ws_summary.cell(r, c1)
        value_cell = ws_summary.cell(r+1, c1)
        title_cell.value = label
        title_cell.font = kpi_title_font
        value_cell.value = val
        value_cell.font = kpi_value_font
        if kind == 'money':
            value_cell.number_format = EXCEL_RUB_FORMAT
        elif kind == 'int':
            value_cell.number_format = EXCEL_INT_FORMAT
        title_cell.alignment = value_cell.alignment = Alignment(horizontal='left', vertical='center')
        for rr in (r, r+1):
            for cc in range(c1, c2+1):
                cell = ws_summary.cell(rr, cc)
                cell.fill = kpi_fill
                cell.border = S['border']

    # Revenue by types
    ws_summary.merge_cells('A13:B13')
    ws_summary['A13'] = 'Выручка по типам станций'
    ws_summary['A13'].font = S['subtitle_font']
    ws_summary['A14'] = 'Тип станции'
    ws_summary['B14'] = 'Выручка'
    S['style_header'](ws_summary, 14)
    station_totals = {'PlayStation': 0, 'Симулятор гонок': 0, 'Nintendo Switch': 0}
    for sale in sales_rows:
        station_totals[sale['stationType']] = station_totals.get(sale['stationType'], 0) + sale['amount']
    row = 15
    for label in ('PlayStation', 'Симулятор гонок', 'Nintendo Switch'):
        ws_summary[f'A{row}'] = label
        _set_summary_value(ws_summary[f'B{row}'], station_totals.get(label, 0), 'money')
        row += 1

    # Payment methods
    ws_summary.merge_cells('A19:B19')
    ws_summary['A19'] = 'Выручка по оплате'
    ws_summary['A19'].font = S['subtitle_font']
    ws_summary['A20'] = 'Оплата'
    ws_summary['B20'] = 'Сумма'
    S['style_header'](ws_summary, 20)
    row = 21
    payment_totals = {'Наличные': 0, 'Карта': 0, 'Перевод': 0}
    for sale in sales_rows:
        payment_totals[sale['payment']] = payment_totals.get(sale['payment'], 0) + sale['amount']
    for label in ('Наличные', 'Карта', 'Перевод'):
        ws_summary[f'A{row}'] = label
        _set_summary_value(ws_summary[f'B{row}'], payment_totals.get(label, 0), 'money')
        row += 1

    # Daily revenue table + heatmap
    ws_summary.merge_cells('A25:B25')
    ws_summary['A25'] = 'Тепловая карта выручки по дням'
    ws_summary['A25'].font = S['subtitle_font']
    daily_header_row = 26
    ws_summary['A26'] = 'Дата'
    ws_summary['B26'] = 'Выручка'
    S['style_header'](ws_summary, 26)
    row = 27
    day_totals = {k: 0 for k in period_keys}
    for sale in sales_rows:
        try:
            iso_key = datetime.strptime(sale['date'], '%d.%m.%Y').strftime('%Y-%m-%d')
        except Exception:
            continue
        day_totals[iso_key] = day_totals.get(iso_key, 0) + sale['amount']
    for key in period_keys:
        ws_summary[f'A{row}'] = _format_date_ru(key)
        _set_summary_value(ws_summary[f'B{row}'], day_totals.get(key, 0), 'money')
        row += 1
    daily_end_row = row - 1
    _apply_daily_heatmap(ws_summary, 2, 27, daily_end_row)

    # Top analytics
    ws_summary['E6'] = 'TOP-аналитика'
    ws_summary['E6'].font = S['subtitle_font']
    ws_summary['E7'] = 'Показатель'
    ws_summary['F7'] = 'Значение'
    for ref in ('E7','F7'):
        ws_summary[ref].fill = S['header_fill']
        ws_summary[ref].font = S['header_font']
        ws_summary[ref].alignment = S['center']
        ws_summary[ref].border = S['border']

    top_items = _top_analytics(session_rows, sales_rows)
    top_rows = [
        (top_items[0][0], top_items[0][1], top_items[0][2]),
        (top_items[1][0], top_items[1][1], top_items[1][2]),
        ('', '', 'blank'),
        (top_items[2][0], top_items[2][1], top_items[2][2]),
        (top_items[3][0], top_items[3][1], top_items[3][2]),
        ('', '', 'blank'),
        (top_items[4][0], top_items[4][1], top_items[4][2]),
        (top_items[5][0], top_items[5][1], top_items[5][2]),
    ]
    top_row = 8
    for label, val, kind in top_rows:
        if kind == 'blank':
            top_row += 1
            continue
        ws_summary.cell(top_row, 5).value = label
        ws_summary.cell(top_row, 5).font = S['bold_font']
        ws_summary.cell(top_row, 5).alignment = S['left']
        _set_summary_value(ws_summary.cell(top_row, 6), val, kind)
        ws_summary.cell(top_row, 6).alignment = S['left']
        top_row += 1

    # Load by stations (compact visual table)
    ws_summary['H6'] = 'Загрузка станций'
    ws_summary['H6'].font = S['subtitle_font']
    ws_summary['H7'] = 'Станция'
    ws_summary['I7'] = 'Часы'
    for ref in ('H7', 'I7'):
        ws_summary[ref].fill = S['header_fill']
        ws_summary[ref].font = S['header_font']
        ws_summary[ref].alignment = S['center']
        ws_summary[ref].border = S['border']
    load_minutes = {}
    for r in session_rows:
        key = r.get('stationName') or '—'
        mins = max(0, int(r.get('_durationMinutes') or 0))
        load_minutes[key] = load_minutes.get(key, 0) + mins
    station_order = ['PS1','PS2','PS3','PS4','Симулятор гонок','Nintendo Switch']
    load_row = 8
    for name in station_order:
        ws_summary.cell(load_row, 8).value = name
        ws_summary.cell(load_row, 8).alignment = S['left']
        hours = round(load_minutes.get(name, 0) / 60, 1)
        ws_summary.cell(load_row, 9).value = hours
        ws_summary.cell(load_row, 9).number_format = '0.0" ч"'
        ws_summary.cell(load_row, 9).alignment = S['left']
        load_row += 1
    load_end_row = load_row - 1
    ws_summary.conditional_formatting.add(
        f'I8:I{load_end_row}',
        DataBarRule(start_type='num', start_value=0, end_type='max', end_value=0, color='5B9BD5', showValue=True)
    )

    # Visual cleanup
    S['style_range'](ws_summary, 13, 23, 1, 2, alignment=S['left'])
    S['style_range'](ws_summary, 7, top_row - 1, 5, 6, alignment=S['left'])
    S['style_range'](ws_summary, 7, load_end_row, 8, 9, alignment=S['left'])
    S['style_range'](ws_summary, 26, daily_end_row, 1, 2, alignment=S['left'])

    for col, width in {'A':18, 'B':20, 'C':16, 'D':16, 'E':28, 'F':34, 'G':12, 'H':28, 'I':16, 'J':20, 'K':16}.items():
        ws_summary.column_dimensions[col].width = width

    from openpyxl.styles import Alignment
    for row in ws_summary.iter_rows(min_row=1, max_row=ws_summary.max_row, min_col=1, max_col=11):
        for cell in row:
            cell.alignment = Alignment(horizontal='left', vertical='center')
    for ref in ('E7','F7','H7','I7','A14','B14','A20','B20','A26','B26'):
        ws_summary[ref].alignment = Alignment(horizontal='left', vertical='center')
    ws_summary.freeze_panes = 'A6'
    _add_revenue_chart(ws_summary, daily_header_row, daily_end_row, anchor='E18')

    # Sessions sheet
    session_table_rows = [
        [r['date'], r['stationName'], r['stationType'], r['paidTime'], r['start'], r['finish'], r['duration'], r['sum'], r['payment'], r['close']]
        for r in session_rows
    ]
    _write_data_sheet(
        ws_sessions,
        headers=['Дата', 'Станция', 'Тип станции', 'Оплачено времени', 'Старт', 'Финиш', 'Длительность', 'Сумма', 'Оплата', 'Завершение'],
        rows=session_table_rows,
        table_name='ReportSessionsTable',
        money_cols=(8,),
        totals={8: 'sum'},
    )

    # Sales sheet
    sales_table_rows = [
        [r['date'], r['time'], r['stationName'], r['stationType'], r['operation'], r['minutes'], r['amount'], r['payment']]
        for r in sales_rows
    ]
    _write_data_sheet(
        ws_sales,
        headers=['Дата', 'Время', 'Станция', 'Тип станции', 'Операция', 'Минуты', 'Сумма', 'Оплата'],
        rows=sales_table_rows,
        table_name='ReportSalesTable',
        money_cols=(7,),
        int_cols=(6,),
        totals={7: 'sum'},
    )

    audit_errors, audit_warnings = _write_audit_sheet(wb, session_rows, sales_rows)

    ws_summary['A33'] = 'Проверка данных'
    ws_summary['A33'].font = S['bold_font']
    ws_summary['A34'] = 'Ошибок не обнаружено' if not (audit_errors or audit_warnings) else f'Ошибки: {audit_errors} | Предупреждения: {audit_warnings}'
    ws_summary['A34'].font = S['bold_font']

    return wb


def _app_dir() -> Path:
    # In PyInstaller EXE: executable dir. In source: script dir.
    if getattr(sys, "frozen", False):
        return Path(sys.executable).resolve().parent
    return Path(__file__).resolve().parent


def _resource_path(*parts: str) -> str:
    """Absolute path to bundled resources (works in source run and PyInstaller EXE)."""
    if getattr(sys, "frozen", False) and hasattr(sys, "_MEIPASS"):
        base = Path(getattr(sys, "_MEIPASS"))  # type: ignore[attr-defined]
    else:
        base = Path(__file__).resolve().parent
    return str(base.joinpath(*parts))


def log(msg: str) -> None:
    try:
        ts = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        (_app_dir() / LOG_FILE).open("a", encoding="utf-8").write(f"[{ts}] {msg}\n")
    except Exception:
        pass


def ensure_single_instance(open_existing: bool = True) -> None:
    """Prevent multiple EXE instances (Windows)."""
    if os.name != "nt":
        return
    try:
        kernel32 = ctypes.windll.kernel32  # type: ignore[attr-defined]
        mutex = kernel32.CreateMutexW(None, True, MUTEX_NAME)
        # Keep a reference so it isn't GC'ed
        globals()["_PSLOUNGE_MUTEX"] = mutex
        last_err = kernel32.GetLastError()
        ERROR_ALREADY_EXISTS = 183
        if last_err == ERROR_ALREADY_EXISTS:
            log("Another instance is already running. Exiting.")
            if open_existing:
                url_path = _app_dir() / URL_TXT
                url = ""
                if url_path.exists():
                    url = url_path.read_text(encoding="utf-8", errors="ignore").strip()

                if url:
                    # Only open the URL if it looks like a running PS Lounge instance.
                    try:
                        import urllib.request
                        health_url = url.rstrip("/") + "/health"
                        with urllib.request.urlopen(health_url, timeout=0.5) as r:
                            body = (r.read(64) or b"").decode("utf-8", errors="ignore")
                        if "OK:PSLOUNGE" in body:
                            webbrowser.open(url)
                        else:
                            log(f"Existing URL does not look like PS Lounge: {health_url} -> {body!r}")
                    except Exception as e:
                        log(f"Failed to verify/open existing URL: {e!r}")
            raise SystemExit(0)
    except SystemExit:
        raise
    except Exception as e:
        # If mutex fails, don't block app start; just log.
        log(f"Single-instance mutex failed: {e!r}")


def port_in_use(host: str, port: int) -> bool:
    """Return True if a TCP server is already listening on host:port.

    Uses connect_ex (reliable on Windows even when another process binds 0.0.0.0:port).
    """
    check_host = "127.0.0.1" if host in ("0.0.0.0", "::") else host
    try:
        with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as s:
            s.settimeout(0.5)
            return s.connect_ex((check_host, port)) == 0
    except Exception:
        # If we cannot check, assume it's in use to be safe.
        return True


def _can_bind(host: str, port: int) -> bool:
    """Best-effort check that we can bind host:port."""
    try:
        with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as s:
            s.setsockopt(socket.SOL_SOCKET, socket.SO_REUSEADDR, 1)
            s.bind((host, port))
        return True
    except OSError:
        return False


def pick_port(host: str, preferred: int) -> int:
    """Prefer preferred; fall back to a free ephemeral port."""
    if not port_in_use(host, preferred) and _can_bind(host, preferred):
        return preferred

    # Ask OS for a free port on the target host
    with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as s:
        s.bind((host, 0))
        return int(s.getsockname()[1])


def write_url_files(url: str) -> None:
    d = _app_dir()
    try:
        (d / URL_TXT).write_text(url, encoding="utf-8")
    except Exception as e:
        log(f"Failed to write {URL_TXT}: {e!r}")

    # Windows .url shortcut
    try:
        content = "[InternetShortcut]\nURL={}\n".format(url)
        (d / URL_SHORTCUT).write_text(content, encoding="utf-8")
    except Exception as e:
        log(f"Failed to write {URL_SHORTCUT}: {e!r}")


def _default_state(source: str = "default") -> dict:
    return {
        "schemaVersion": SCHEMA_VERSION,
        "version": 1,
        "lastModified": 0,
        "stations": None,
        "sessions": None,
        "settings": None,
        "source": source,
    }


def _is_valid_station_list(value: object) -> bool:
    if not isinstance(value, list) or not value:
        return False
    allowed_status = {"idle", "running", "grace", "overdue"}
    seen_ids: set[int] = set()
    for idx, item in enumerate(value, start=1):
        if not isinstance(item, dict):
            return False
        sid = _safe_int(item.get("id"), idx)
        if sid <= 0 or sid in seen_ids:
            return False
        seen_ids.add(sid)
        if item.get("status") not in allowed_status:
            return False
        if item.get("startTime") is not None and not isinstance(item.get("startTime"), (int, float)):
            return False
        if item.get("endTime") is not None and not isinstance(item.get("endTime"), (int, float)):
            return False
        if item.get("stationType") is not None and _normalize_station_type_key(item.get("stationType"), "") == "":
            return False
    return True


def _is_valid_sessions_map(value: object) -> bool:
    if value is None:
        return True
    if not isinstance(value, dict):
        return False
    for day_key, entries in value.items():
        if not isinstance(day_key, str) or not isinstance(entries, list):
            return False
        try:
            datetime.strptime(day_key, "%Y-%m-%d")
        except Exception:
            return False
        for item in entries:
            if not _is_valid_session_record(item):
                return False
    return True


def _is_valid_settings_blob(value: object) -> bool:
    if value is None:
        return True
    if not isinstance(value, dict):
        return False
    defs = value.get("stationDefinitions")
    if defs is not None:
        if not isinstance(defs, list) or not defs:
            return False
        seen_ids: set[int] = set()
        for idx, item in enumerate(defs):
            normalized = _normalize_station_definition(item, idx)
            if normalized is None or normalized["id"] in seen_ids:
                return False
            seen_ids.add(normalized["id"])
    return True


def _is_valid_state_payload(data: object) -> bool:
    if not isinstance(data, dict):
        return False
    if not isinstance(data.get("lastModified", 0), int):
        return False
    if not _is_valid_station_list(data.get("stations")):
        return False
    if not _is_valid_sessions_map(data.get("sessions")):
        return False
    if not _is_valid_settings_blob(data.get("settings")):
        return False
    return True


def _trim_history_dir(history_dir: Path, retention_days: int) -> None:
    if retention_days <= 0 or not history_dir.exists():
        return
    cutoff = datetime.now().date() - timedelta(days=retention_days)
    for item in history_dir.glob("*.json"):
        try:
            d = datetime.strptime(item.stem, "%Y-%m-%d").date()
        except Exception:
            continue
        if d < cutoff:
            try:
                item.unlink(missing_ok=True)
            except Exception as e:
                log(f"Failed to trim history snapshot {item.name}: {e!r}")


def _read_json_candidate(path: Path) -> dict | None:
    try:
        if not path.exists():
            return None
        data = json.loads(path.read_text(encoding="utf-8"))
        if not isinstance(data, dict):
            return None
        if not _is_valid_state_payload(data):
            log(f"State candidate {path.name} failed validation")
            return None
        data.setdefault("schemaVersion", SCHEMA_VERSION)
        data.setdefault("version", 1)
        data.setdefault("lastModified", 0)
        return data
    except Exception as e:
        log(f"Failed to read json candidate {path.name}: {e!r}")
        return None


def _read_state_file() -> dict:
    app_dir = _app_dir()
    primary = app_dir / STATE_FILE
    backup = app_dir / STATE_BAK_FILE

    primary_data = _read_json_candidate(primary)
    if primary_data is not None:
        primary_data["source"] = "primary"
        return primary_data

    backup_data = _read_json_candidate(backup)
    if backup_data is not None:
        backup_data["source"] = "backup"
        return backup_data

    return _default_state()


def _atomic_write_json(path: Path, data: dict) -> None:
    tmp = path.with_suffix(path.suffix + ".tmp")
    payload = json.dumps(data, ensure_ascii=False)
    tmp.write_text(payload, encoding="utf-8")
    json.loads(tmp.read_text(encoding="utf-8"))

    backup_path = _app_dir() / STATE_BAK_FILE
    if path.exists():
        try:
            shutil.copy2(path, backup_path)
        except Exception as e:
            log(f"Failed to refresh backup file: {e!r}")

    tmp.replace(path)

    try:
        history_dir = _app_dir() / STATE_HISTORY_DIR
        history_dir.mkdir(exist_ok=True)
        date_key = datetime.now().strftime("%Y-%m-%d")
        (history_dir / f"{date_key}.json").write_text(payload, encoding="utf-8")
        _trim_history_dir(history_dir, int(os.environ.get("PS_LOUNGE_HISTORY_DAYS", str(STATE_HISTORY_RETENTION_DAYS))))
    except Exception as e:
        log(f"Failed to write daily state snapshot: {e!r}")

def _trim_sessions(sessions: dict, retention_days: int) -> dict:
    if retention_days <= 0:
        return sessions
    cutoff = datetime.now().date() - timedelta(days=retention_days)
    out = {}
    for k, v in sessions.items():
        try:
            d = datetime.strptime(k, "%Y-%m-%d").date()
        except Exception:
            # keep unknown keys
            out[k] = v
            continue
        if d >= cutoff:
            out[k] = v
    return out


# Flask app
app = Flask(
    __name__,
    static_folder=_resource_path("static"),
    template_folder=_resource_path("templates"),
    static_url_path="/static",
)


def _asset_version(relative_path: str) -> str:
    try:
        return str(int(Path(_resource_path(relative_path)).stat().st_mtime))
    except Exception:
        return "1"


@app.context_processor
def inject_asset_helpers():
    def asset_url(filename: str) -> str:
        return url_for("static", filename=filename, v=_asset_version(f"static/{filename}"))

    return {"asset_url": asset_url}


@app.get("/")
def index():
    return render_template("index.html")


@app.get("/health")
def health():
    # Used by the launcher to verify that the running instance is PS Lounge.
    return "OK:PSLOUNGE", 200


@app.get("/api/license/status")
def api_license_status():
    return jsonify(_license_status())


@app.post("/api/license/activate")
def api_license_activate():
    payload = request.get_json(silent=True) or {}
    token = str(payload.get("token") or payload.get("key") or "").strip()
    if not token:
        return jsonify({"ok": False, "error": "missing_token"}), 400
    try:
        status = _activate_license_token(token)
        return jsonify(status)
    except ValueError as e:
        return jsonify({"ok": False, "error": "invalid_license", "detail": str(e)}), 400
    except Exception as e:
        log(f"Failed to activate license: {e!r}")
        return jsonify({"ok": False, "error": "license_write_failed"}), 500


@app.get("/api/backup")
def api_get_backup():
    denied = _require_license()
    if denied:
        return denied
    return jsonify(_read_state_file())


@app.post("/api/backup")
def api_set_backup():
    denied = _require_license()
    if denied:
        return denied
    payload = request.get_json(silent=True) or {}
    if not isinstance(payload, dict):
        return jsonify({"ok": False, "error": "bad json"}), 400

    last_modified = payload.get("lastModified")
    stations = payload.get("stations")
    sessions = payload.get("sessions")
    settings = payload.get("settings")

    if not isinstance(last_modified, int):
        return jsonify({"ok": False, "error": "lastModified must be int"}), 400

    # Read current to prevent older overwrite
    current = _read_state_file()
    cur_lm = int(current.get("lastModified") or 0)
    if last_modified < cur_lm:
        return jsonify({"ok": True, "skipped": True, "reason": "older_than_current", "currentLastModified": cur_lm})

    # Trim sessions growth
    retention_days = int(os.environ.get("PS_LOUNGE_RETENTION_DAYS", "60"))
    if isinstance(sessions, dict):
        sessions = _trim_sessions(sessions, retention_days)

    data = {
        "schemaVersion": SCHEMA_VERSION,
        "version": 1,
        "lastModified": last_modified,
        "stations": stations,
        "sessions": sessions,
        "settings": settings,
    }

    if not _is_valid_state_payload(data):
        return jsonify({"ok": False, "error": "invalid_state_shape"}), 400

    try:
        _atomic_write_json(_app_dir() / STATE_FILE, data)
        return jsonify({"ok": True})
    except Exception as e:
        log(f"Failed to write state file: {e!r}")
        return jsonify({"ok": False, "error": "write_failed"}), 500





@app.get("/api/export/today.xlsx")
def api_export_today_xlsx():
    denied = _require_license()
    if denied:
        return denied
    try:
        key = _parse_date_key(request.args.get("date"), datetime.now().strftime("%Y-%m-%d"))
    except ValueError:
        return jsonify({"ok": False, "error": "invalid_date"}), 400
    state = _read_state_file()
    session_rows, sales_rows, period_keys = _collect_report_rows(state, key, key)
    wb = _build_shift_workbook(key, session_rows, sales_rows, state)
    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)
    filename = f"PS_Lounge_Shift_{datetime.strptime(key, '%Y-%m-%d').strftime('%d-%m-%Y')}.xlsx"
    return send_file(
        bio,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=filename,
    )


@app.get("/api/export/report.xlsx")
def api_export_report_xlsx():
    denied = _require_license()
    if denied:
        return denied
    try:
        from_key = _parse_date_key(request.args.get("from"), datetime.now().strftime("%Y-%m-%d"))
        to_key = _parse_date_key(request.args.get("to"), from_key)
    except ValueError:
        return jsonify({"ok": False, "error": "invalid_date"}), 400
    station_type = _normalize_station_type_key(request.args.get("stationType"), "")
    payment = request.args.get("payment") or ""

    if from_key > to_key:
        from_key, to_key = to_key, from_key

    state = _read_state_file()
    session_rows, sales_rows, period_keys = _collect_report_rows(state, from_key, to_key, station_type, payment)
    wb = _build_report_workbook(from_key, to_key, session_rows, sales_rows, period_keys, state)
    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)

    if from_key == to_key:
        filename = f"PS_Lounge_Report_{datetime.strptime(from_key, '%Y-%m-%d').strftime('%d-%m-%Y')}.xlsx"
    else:
        filename = f"PS_Lounge_Report_{datetime.strptime(from_key, '%Y-%m-%d').strftime('%d-%m-%Y')}_{datetime.strptime(to_key, '%Y-%m-%d').strftime('%d-%m-%Y')}.xlsx"
    return send_file(
        bio,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=filename,
    )


_orig_build_shift_workbook = _build_shift_workbook
_orig_build_report_workbook = _build_report_workbook


def _restyle_station_rows(ws, start_row: int, names: list[str], totals: dict[str, int], value_col: str, old_count: int = 6) -> None:
    target_count = max(1, len(names))
    delta = target_count - old_count
    insert_at = start_row + old_count
    if delta > 0:
        ws.insert_rows(insert_at, delta)
    elif delta < 0:
        ws.delete_rows(start_row + target_count, -delta)
    for offset, name in enumerate(names or ['—']):
        row = start_row + offset
        ws[f'A{row}' if value_col == 'B' else f'H{row}'] = name
        cell = ws[f'{value_col}{row}']
        cell.value = totals.get(name, 0)
        if value_col == 'B':
            cell.number_format = EXCEL_RUB_FORMAT
        else:
            cell.number_format = '0.0" ч"'


def _build_shift_workbook(day_key: str, session_rows: list[dict], sales_rows: list[dict], state: dict | None = None):
    wb = _orig_build_shift_workbook(day_key, session_rows, sales_rows)
    names = _report_station_names(state, session_rows, sales_rows)
    totals = {name: 0 for name in names}
    for sale in sales_rows:
        name = str(sale.get('stationName') or '').strip()
        if name:
            totals[name] = totals.get(name, 0) + _money_int(sale.get('amount'))
    _restyle_station_rows(wb.active, 13, names, totals, 'B')
    return wb


def _build_report_workbook(from_key: str, to_key: str, session_rows: list[dict], sales_rows: list[dict], period_keys: list[str], state: dict | None = None):
    wb = _orig_build_report_workbook(from_key, to_key, session_rows, sales_rows, period_keys)
    ws = wb.worksheets[0]
    names = _report_station_names(state, session_rows, sales_rows)
    load_minutes: dict[str, int] = {name: 0 for name in names}
    for row in session_rows:
        name = str(row.get('stationName') or '').strip()
        if name:
            load_minutes[name] = load_minutes.get(name, 0) + max(0, _safe_int(row.get('_durationMinutes'), 0))
    load_hours = {name: round(load_minutes.get(name, 0) / 60, 1) for name in names}
    _restyle_station_rows(ws, 8, names, load_hours, 'I')
    return wb


def main() -> None:
    ensure_single_instance(open_existing=True)

    host = os.environ.get("PS_LOUNGE_HOST", "127.0.0.1")
    preferred_port = int(os.environ.get("PS_LOUNGE_PORT", "5000"))
    debug = os.environ.get("PS_LOUNGE_DEBUG", "0") == "1"
    no_browser = os.environ.get("PS_LOUNGE_NO_BROWSER", "0") == "1"

    if port_in_use(host, preferred_port) or not _can_bind(host, preferred_port):
        log(f"Preferred port {preferred_port} is busy; selecting a free port")
    port = pick_port(host, preferred_port)
    url = f"http://{host}:{port}/"

    log(f"Starting {APP_NAME} on {url} (preferred {preferred_port})")
    write_url_files(url)

    # Open browser once on start.
    if not no_browser:
        if os.environ.get("WERKZEUG_RUN_MAIN") != "true":
            threading.Thread(
                target=lambda: (time.sleep(0.6), webbrowser.open(url)),
                daemon=True
            ).start()

    # Never use reloader in packaged app: it breaks single-instance behavior.
    app.run(host=host, port=port, debug=debug, use_reloader=False)


if __name__ == "__main__":
    main()
